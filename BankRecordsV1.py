#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
温岭纪委初核工具（单文件重构优化版）

"""

from __future__ import annotations

import builtins
import datetime as _dt
import re
import sys
import unicodedata
import threading
import warnings
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from chinese_calendar import is_holiday, is_workday
try:
    from chinese_calendar import get_holiday_detail, Holiday
except Exception:
    get_holiday_detail, Holiday = None, None

try:
    from lunardate import LunarDate
except Exception:
    LunarDate = None

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# ========= 输出目录配置 =========
OUTPUT_ROOT_NAME = "分析结果"

# ========= 常量 / 正则 =========
SKIP_HEADER_KEYWORDS = ["反洗钱-电子账户交易明细", "信用卡消费明细"]
FULL_TS_PAT = re.compile(r"\d{4}-\d{2}-\d{2}-\d{2}\.\d{2}\.\d{2}\.\d+")
COMPACT_DT_DIGITS_RE = re.compile(r"^\d{12,16}$")
ONLY_DIGITS_RE = re.compile(r"\D+")
_MOBILE_PAT = re.compile(r"(?:\+?86[-\s]?)?(1[3-9]\d{9})")

CH_NAME_2_4_RE = re.compile(r"^[\u4e00-\u9fff]{2,4}$")  # 交易对手Top8仅姓名
HOTEL_SUFFIX, FLIGHT_SUFFIX = "旅馆同住", "同机"
HOTEL_COL_NAME, FLIGHT_COL_NAME = "姓名", "中文名"

COMM_SUFFIX = "通信"
COMM_FILE_RE = re.compile(rf"^(.+?)-{COMM_SUFFIX}$")

WORK_START_HOUR, WORK_END_HOUR = 9, 18
NIGHT_START, NIGHT_END = 23, 5
FESTIVAL_NAMES = ["春节", "中秋节", "端午节", "七夕节", "5月20日"]

TEMPLATE_COLS = [
    "序号","查询对象","反馈单位","查询项","查询账户","查询卡号","交易类型","借贷标志","币种","交易金额","账户余额","交易时间","交易流水号",
    "本方账号","本方卡号","交易对方姓名","交易对方账户","交易对方卡号","交易对方证件号码","交易对手余额","交易对方账号开户行","交易摘要",
    "交易网点名称","交易网点代码","日志号","传票号","凭证种类","凭证号","现金标志","终端号","交易是否成功","交易发生地","商户名称","商户号",
    "IP地址","MAC","交易柜员号","备注",
]

STRICT_CONTACTS_REQUIRED = ["姓名", "职务", "号码"]
NAME_CANDIDATE_COLS = ["账户名称","户名","账户名","账号名称","账号名","姓名","客户名称","查询对象"]


# ========= 状态 =========
@dataclass
class AppState:
    OUT_DIR: Optional[Path] = None

    CONTACT_PHONE_TO_NAME_TITLE: Dict[str, Tuple[str, str]] = field(default_factory=dict)
    CALLLOG_NAME_TO_TITLE: Dict[str, str] = field(default_factory=dict)

    COMM_STATS_ALL: Optional[pd.DataFrame] = None
    COMM_STATS_BY_PERSON: Dict[str, pd.DataFrame] = field(default_factory=dict)
    COMM_RECORDS_BY_PERSON: Dict[str, int] = field(default_factory=dict)

    COMM_FILES_COUNT: int = 0
    COMM_RECORDS_COUNT: int = 0

    RAW_TXN_COUNT: int = 0
    DUP_TXN_COUNT: int = 0
    FINAL_TXN_COUNT: int = 0

    CLUE_SUMMARY_BY_PERSON: Dict[str, pd.DataFrame] = field(default_factory=dict)
    CLUE_DETAIL_BY_PERSON: Dict[str, pd.DataFrame] = field(default_factory=dict)

    def in_out_dir(self, p: Path) -> bool:
        if self.OUT_DIR is None:
            return False
        try:
            return p.resolve().is_relative_to(self.OUT_DIR.resolve())
        except AttributeError:
            return str(p.resolve()).startswith(str(self.OUT_DIR.resolve()))

STATE = AppState()


# ========= 输出目录 =========
def out_root() -> Path:
    if STATE.OUT_DIR is None:
        raise RuntimeError("OUT_DIR 未初始化")
    return STATE.OUT_DIR


def _ensure_dir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def _sanitize_person_name(name: str) -> str:
    name = (name or "").strip() or "未知"
    name = re.sub(r"[\\/:*?\"<>|\r\n\t]+", "_", name)
    return name[:40]


def out_person_root(person: str) -> Path:
    return _ensure_dir(out_root() / _sanitize_person_name(person))


def _person_subdir(person: str, dirname: str) -> Path:
    return _ensure_dir(out_person_root(person) / dirname)


def out_bank_dir(person: str) -> Path:
    return _person_subdir(person, "银行")


def out_comm_dir(person: str) -> Path:
    return _person_subdir(person, "通信")


def out_report_dir(person: str) -> Path:
    return _person_subdir(person, "报告")


def out_gongan_dir(person: str) -> Path:
    return _person_subdir(person, "公安信息")


def out_comm_summary_dir() -> Path:
    return _ensure_dir(out_root() / "通信汇总")


# ========= 通用工具 =========
def safe_str(x: Any) -> str:
    if x is None:
        return ""
    try:
        if pd.isna(x):
            return ""
    except Exception:
        pass
    s = str(x)
    return "" if s.lower() == "nan" else s


def clean_columns(df: pd.DataFrame) -> pd.DataFrame:
    """返回列名已转字符串并去首尾空格的 DataFrame 副本。"""
    d = df.copy()
    d.columns = pd.Index(d.columns).astype(str).str.strip()
    return d

def _parse_compact_datetime(s: Any) -> Optional[str]:
    raw = safe_str(s).strip()
    if not raw:
        return None
    digits = ONLY_DIGITS_RE.sub("", raw)
    if not COMPACT_DT_DIGITS_RE.fullmatch(digits):
        return None
    try:
        if len(digits) >= 14:
            y,m,d,hh,mm,ss = map(int, [digits[0:4],digits[4:6],digits[6:8],digits[8:10],digits[10:12],digits[12:14]])
        else:
            y,m,d,hh,mm,ss = map(int, [digits[0:4],digits[4:6],digits[6:8],digits[8:10],digits[10:12],0])
        return _dt.datetime(y,m,d,hh,mm,ss).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return None


def normalize_phone_cell(x: Any) -> str:
    if x is None or (isinstance(x, float) and not np.isfinite(x)):
        return ""
    if isinstance(x, (int, np.integer)):
        s = str(int(x)); m = _MOBILE_PAT.search(s); return m.group(1) if m else s
    if isinstance(x, (float, np.floating)):
        try:
            s = str(int(x)); m = _MOBILE_PAT.search(s); return m.group(1) if m else s
        except Exception:
            pass
    s = safe_str(x).strip().replace("\u00A0", " ")
    if not s:
        return ""
    m = _MOBILE_PAT.search(s)
    if m:
        return m.group(1)
    if re.fullmatch(r"\d+(\.0+)?", s):
        return s.split(".")[0]
    if re.fullmatch(r"[0-9]+(\.[0-9]+)?[eE][+-]?[0-9]+", s):
        try:
            d = Decimal(s).quantize(Decimal(1))
            ss = format(d, "f")
            m2 = _MOBILE_PAT.search(ss)
            return m2.group(1) if m2 else ss
        except InvalidOperation:
            pass
    only_digits = re.sub(r"\D", "", s)
    if len(only_digits) >= 11:
        m3 = _MOBILE_PAT.search(only_digits)
        if m3:
            return m3.group(1)
    return only_digits

def normalize_phone_series(s: pd.Series) -> pd.Series:
    if s is None or len(s) == 0:
        return pd.Series([], dtype=object)
    ss = s.astype(str).str.replace("\u00A0", " ", regex=False).str.strip()
    sci_like = ss.str.fullmatch(r"[0-9]+(\.[0-9]+)?([eE][+-]?[0-9]+)?")
    if sci_like.any():
        def _fix(x: str) -> str:
            try:
                if re.fullmatch(r"\d+\.0+", x): return x.split(".")[0]
                if re.fullmatch(r"[0-9]+(\.[0-9]+)?([eE][+-]?[0-9]+)?", x): return str(int(float(x)))
            except Exception:
                pass
            return x
        ss.loc[sci_like] = ss.loc[sci_like].map(_fix)
    extracted = ss.str.extract(_MOBILE_PAT, expand=False)
    miss = extracted.isna()
    if miss.any():
        extracted.loc[miss] = ss.loc[miss].str.replace(r"\D", "", regex=True)
    long_mask = extracted.str.len().fillna(0) >= 11
    if long_mask.any():
        extracted.loc[long_mask] = extracted.loc[long_mask].map(lambda x: (_MOBILE_PAT.search(x).group(1) if x and _MOBILE_PAT.search(x) else x) if x else "")
    return extracted.fillna("")


def save_df_auto_width(
    df: pd.DataFrame, filename: Path | str, sheet_name: str = "Sheet1", index: bool = False,
    engine: str = "xlsxwriter", min_width: int = 6, max_width: int = 50
) -> Path:
    filename = Path(filename)
    if not filename.is_absolute():
        filename = out_root() / filename
    filename = filename.with_suffix(".xlsx")
    filename.parent.mkdir(parents=True, exist_ok=True)
    if filename.exists():
        try: filename.unlink()
        except Exception: pass

    df = df.copy().replace({np.nan: ""})

    if engine == "xlsxwriter":
        with pd.ExcelWriter(filename, engine="xlsxwriter", mode="w") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=index)
            ws = writer.sheets[sheet_name]
            for i, col in enumerate(df.columns):
                s = df[col].astype(str)
                width = max(min_width, min(max(s.map(len).max(), len(str(col))) + 2, max_width))
                ws.set_column(i, i, width)
    else:
        with pd.ExcelWriter(filename, engine="openpyxl", mode="w") as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=index)
        from openpyxl import load_workbook
        wb = load_workbook(filename); ws = wb[sheet_name]
        for col_cells in ws.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells) + 2
            ws.column_dimensions[col_cells[0].column_letter].width = max(min_width, min(width, max_width)) + 5
        wb.save(filename)
    return filename


def save_excel_sheets_auto_width(
    sheets: Dict[str, pd.DataFrame], filename: Path | str, index: bool = False,
    min_width: int = 6, max_width: int = 60
) -> Path:
    """按多 sheet 导出 Excel，并做基础列宽适配。"""
    filename = Path(filename)
    if not filename.is_absolute():
        filename = out_root() / filename
    filename = filename.with_suffix(".xlsx")
    filename.parent.mkdir(parents=True, exist_ok=True)
    if filename.exists():
        try:
            filename.unlink()
        except Exception:
            pass

    clean_sheets: Dict[str, pd.DataFrame] = {}
    for name, df in sheets.items():
        sheet = re.sub(r"[\\/:*?\[\]]+", "_", safe_str(name).strip() or "Sheet")[:31]
        d = df.copy() if df is not None else pd.DataFrame()
        clean_sheets[sheet] = d.replace({np.nan: ""})

    with pd.ExcelWriter(filename, engine="openpyxl", mode="w") as writer:
        for sheet, df in clean_sheets.items():
            df.to_excel(writer, sheet_name=sheet, index=index)

    from openpyxl import load_workbook
    wb = load_workbook(filename)
    for ws in wb.worksheets:
        for col_cells in ws.columns:
            values = [str(c.value) if c.value is not None else "" for c in col_cells]
            width = max((len(v) for v in values), default=0) + 2
            ws.column_dimensions[col_cells[0].column_letter].width = max(min_width, min(width, max_width))
    wb.save(filename)
    return filename


# ========= 文件/人名推断 =========


def _person_from_comm_filename(p: Path) -> str:
    m = COMM_FILE_RE.match(p.stem)
    return safe_str(m.group(1)).strip() if m else ""

# ========= 表头定位/读取 =========


# ========= 同住/同机（职务匹配） =========
def _find_header_row_contains(xls: pd.ExcelFile, sheet_name: str, required_cols: List[str], scan_rows: int = 60) -> Optional[int]:
    try:
        df0 = xls.parse(sheet_name, header=None, nrows=scan_rows)
    except Exception:
        return None
    req = {safe_str(c).strip() for c in required_cols if safe_str(c).strip()}
    if not req:
        return 0
    for i, row in df0.iterrows():
        vals = {safe_str(v).strip() for v in row.values}
        if req.issubset(vals):
            return i
    return None

def _infer_person_from_special_filename(p: Path) -> str:
    m = re.match(r"^(.+?)-(旅馆同住|同机)$", p.stem)
    if m:
        return safe_str(m.group(1)).strip()
    if "-" in p.stem:
        return safe_str(p.stem.split("-", 1)[0]).strip()
    return ""

def _insert_title_column_after(df: pd.DataFrame, base_col: str, title_col: str, title_map: Dict[str, str]) -> pd.DataFrame:
    if df is None or df.empty:
        return df
    d = clean_columns(df)
    if base_col not in d.columns:
        return d
    names = d[base_col].map(safe_str).str.strip()
    titles = names.map(title_map).fillna("")
    if title_col in d.columns:
        d[title_col] = titles
        cols = list(d.columns); cols.remove(title_col)
        idx = cols.index(base_col) + 1
        return d[cols[:idx] + [title_col] + cols[idx:]]
    insert_at = list(d.columns).index(base_col) + 1
    d.insert(insert_at, title_col, titles)
    return d

def annotate_hotel_and_flight_files(root: Path, name_to_title: Dict[str, str]) -> None:
    if not name_to_title:
        print("ℹ️ 未生成通信姓名->职务映射，跳过“旅馆同住/同机”职务匹配。")
        return
    all_xls = [p for p in root.rglob("*.xls*") if (not STATE.in_out_dir(p)) and (not p.name.startswith("~$"))]
    targets = [p for p in all_xls if (p.stem.endswith(f"-{HOTEL_SUFFIX}") or p.stem.endswith(f"-{FLIGHT_SUFFIX}"))]
    if not targets:
        print("ℹ️ 未发现“(姓名)-旅馆同住”或“(姓名)-同机”文件，跳过职务匹配。")
        return
    print(f"🧩 检测到 {len(targets)} 个“旅馆同住/同机”文件，开始按通信映射补全职务...")

    for p in targets:
        person = safe_str(_infer_person_from_special_filename(p)).strip() or "未知"
        kind = HOTEL_SUFFIX if p.stem.endswith(f"-{HOTEL_SUFFIX}") else FLIGHT_SUFFIX
        base_col = HOTEL_COL_NAME if kind == HOTEL_SUFFIX else FLIGHT_COL_NAME
        try:
            xls = pd.ExcelFile(p)
        except Exception as e:
            print("❌ 同住/同机文件载入失败", p.name, e)
            continue

        frames: List[pd.DataFrame] = []
        for sht in xls.sheet_names:
            try:
                hdr = _find_header_row_contains(xls, sht, [base_col], scan_rows=80)
                if hdr is None:
                    print(f"  • 跳过 {p.name}/{sht}：未找到表头（需要列：{base_col}）")
                    continue
                df0 = clean_columns(xls.parse(sheet_name=sht, header=hdr))
                if base_col not in df0.columns:
                    print(f"  • 跳过 {p.name}/{sht}：解析后仍缺少列：{base_col}")
                    continue
                df1 = _insert_title_column_after(df0, base_col, "职务", name_to_title)
                df1.insert(0, "__来源sheet__", sht)
                frames.append(df1)
            except Exception as e:
                print("❌ 同住/同机解析失败", f"{p.name}->{sht}", e)

        if not frames:
            print(f"ℹ️ {p.name} 未生成可输出的数据（可能缺少列：{base_col}）")
            continue

        merged = pd.concat(frames, ignore_index=True).replace({np.nan: ""})
        out_dir = out_gongan_dir(person)
        out_file = out_dir / f"{p.stem}-已标注"
        save_df_auto_width(merged, out_file, index=False, engine="openpyxl")
        print(f"✅ {kind}职务匹配导出：{out_dir / (p.stem + '-已标注.xlsx')}")

# ========= 通讯录读取（严格列名：姓名/职务/号码） =========
def _iter_builtin_contacts_files() -> List[Path]:
    candidates = ["公职人员-内置通讯录.xlsx", "公职人员-内置通讯录.xls","社保-内置通讯录.xlsx", "社保-内置通讯录.xls"]
    base_dirs: List[Path] = []
    try: base_dirs.append(Path(__file__).parent.resolve())
    except Exception: pass
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        base_dirs.append(Path(sys._MEIPASS).resolve())
    try: base_dirs.append(Path.cwd().resolve())
    except Exception: pass

    out, seen = [], set()
    for b in base_dirs:
        for name in candidates:
            p = b / name
            try: rp = str(p.resolve())
            except Exception: rp = str(p)
            if p.exists() and rp not in seen:
                out.append(p); seen.add(rp)
    return out

def _guess_header_row_strict(xls: pd.ExcelFile, sheet_name: str, scan_rows: int = 30) -> Optional[int]:
    df0 = xls.parse(sheet_name, header=None, nrows=scan_rows)
    req = set(STRICT_CONTACTS_REQUIRED)
    for i, row in df0.iterrows():
        vals = {safe_str(v).strip() for v in row.values}
        if req.issubset(vals):
            return i
    return None

def load_contacts_phone_map_strict(root: Path) -> Dict[str, Tuple[str, str]]:
    print("正在读取通讯录（列名）......")
    builtin_files = _iter_builtin_contacts_files()
    for bp in builtin_files:
        print(f"  • 使用内置通讯录：{bp.name}")

    repo_files = [p for p in root.rglob("*通讯录*.xls*") if ("已标注" not in p.stem) and (not STATE.in_out_dir(p))]
    all_files, seen = [], set()
    for p in [*builtin_files, *repo_files]:
        try:
            rp = str(p.resolve())
        except Exception:
            rp = str(p)
        if rp not in seen:
            all_files.append(p)
            seen.add(rp)

    if not all_files:
        print("ℹ️ 未发现可用的通讯录。")
        return {}

    merged_name: Dict[str, str] = {}
    merged_titles: Dict[str, List[str]] = {}

    for p in all_files:
        try:
            xls = pd.ExcelFile(p)
        except Exception as e:
            print("❌ 通讯录载入失败", p.name, e)
            continue

        for sht in xls.sheet_names:
            try:
                hdr = _guess_header_row_strict(xls, sht, 40)
                if hdr is None:
                    print(f"  • 跳过 {p.name}/{sht}：未找到表头（需要：姓名/职务/号码）")
                    continue
                df = clean_columns(xls.parse(sht, header=hdr))
                if not set(STRICT_CONTACTS_REQUIRED).issubset(set(df.columns)):
                    print(f"  • 跳过 {p.name}/{sht}：缺少列 {STRICT_CONTACTS_REQUIRED}")
                    continue

                nm = df["姓名"].astype(str).str.strip()
                tt = df["职务"].astype(str).str.strip()
                ph = normalize_phone_series(df["号码"]).astype(str).str.strip()

                dtmp = pd.DataFrame({"号码": ph, "姓名": nm, "职务": tt})
                dtmp = dtmp[dtmp["号码"] != ""]
                before = len(dtmp)
                uniq_cnt = int(dtmp["号码"].nunique())

                print(f"  • 通讯录 {p.name}/{sht}：载入 {len(df)} 行，命中号码 {uniq_cnt}（去空前 {before}）")

                for _, r in dtmp.iterrows():
                    phone = safe_str(r["号码"]).strip()
                    if not phone:
                        continue

                    name = safe_str(r.get("姓名", "")).strip()
                    title = safe_str(r.get("职务", "")).strip()

                    if name and not merged_name.get(phone):
                        merged_name[phone] = name

                    if title:
                        lst = merged_titles.setdefault(phone, [])
                        if title not in lst:
                            lst.append(title)

            except Exception as e:
                print("❌ 通讯录解析失败", f"{p.name}->{sht}", e)

    all_phones = set(merged_name) | set(merged_titles)
    merged: Dict[str, Tuple[str, str]] = {}
    for phone in all_phones:
        name = merged_name.get(phone, "")
        titles = "、".join(merged_titles.get(phone, []))
        merged[phone] = (name, titles)

    print(f"✅ 通讯录号码映射加载完成：{len(merged)} 条。")
    return merged


# ========= 节日/时间特征 =========
def _flag_offwork(ts: pd.Series) -> pd.Series:
    h = ts.dt.hour
    return (h < WORK_START_HOUR) | (h >= WORK_END_HOUR)

def _flag_late_night(ts: pd.Series) -> pd.Series:
    h = ts.dt.hour
    return (h >= NIGHT_START) | (h < NIGHT_END)

def _is_festival_day_lunar(g_date: _dt.date) -> str:
    if g_date.month == 5 and g_date.day == 20:
        return "5月20日"

    if LunarDate is not None:
        try:
            ld = LunarDate.fromSolarDate(g_date.year, g_date.month, g_date.day)  # type: ignore
            m, d = ld.month, ld.day
            if m == 1 and 1 <= d <= 15: return "春节"
            if m == 8 and d == 15: return "中秋节"
            if m == 5 and d == 5: return "端午节"
            if m == 7 and d == 7: return "七夕节"
        except Exception:
            pass

    if get_holiday_detail is not None:
        try:
            is_hol, hol = get_holiday_detail(g_date)
            if is_hol and hol is not None:
                name = getattr(hol, "name", str(hol))
                if (Holiday is not None and hol == Holiday.SpringFestival) or "SpringFestival" in name or "春节" in name: return "春节"
                if (Holiday is not None and hol == Holiday.MidAutumnFestival) or "MidAutumn" in name or "中秋" in name: return "中秋节"
                if (Holiday is not None and hol == Holiday.DragonBoatFestival) or "DragonBoat" in name or "端午" in name: return "端午节"
        except Exception:
            pass

    return ""

def _festival_series(ts: pd.Series) -> pd.Series:
    res = pd.Series([""] * len(ts), index=ts.index, dtype=object)
    idx = ts.notna()
    if not idx.any():
        return res
    dates = ts[idx].dt.date
    res.loc[idx] = [_is_festival_day_lunar(d) for d in dates]
    return res


# ========= 通信标注 + 统计 =========
def _find_header_row_exact(xls: pd.ExcelFile, sheet_name: str, required_cols: List[str], scan_rows: int = 40) -> Optional[int]:
    df0 = xls.parse(sheet_name, header=None, nrows=scan_rows)
    req = set(required_cols)
    for i, row in df0.iterrows():
        vals = {safe_str(v).strip() for v in row.values}
        if req.issubset(vals):
            return i
    return None

def _compose_datetime_from_cols_relaxed(df: pd.DataFrame) -> pd.Series:
    if "通话时间" in df.columns:
        ser_raw = df["通话时间"].map(safe_str).str.strip()
        ser_dt = ser_raw.map(lambda s: _parse_compact_datetime(s) or s)
        ser = pd.to_datetime(ser_dt, errors="coerce")
        if ser.notna().any():
            return ser

    c_date = next((c for c in ["日期","发生日期","通话日期"] if c in df.columns), None)
    c_time = next((c for c in ["时间","发生时间","通话时间","开始时间","呼叫时间"] if c in df.columns), None)
    if c_date and c_time:
        combo = df[c_date].map(safe_str).str.strip() + " " + df[c_time].map(safe_str).str.strip()
        return pd.to_datetime(combo.map(lambda s: _parse_compact_datetime(s) or s), errors="coerce")
    if c_date:
        return pd.to_datetime(df[c_date], errors="coerce")
    return pd.to_datetime(pd.Series([pd.NaT] * len(df), index=df.index), errors="coerce")

def _parse_duration_to_seconds(x: Any) -> float:
    s = safe_str(x).strip()
    if not s:
        return np.nan
    if re.fullmatch(r"\d+(\.\d+)?([eE][+-]?\d+)?", s):
        try: return float(s)
        except Exception: pass
    if ":" in s:
        parts = s.split(":")
        try:
            parts = [int(float(p)) for p in parts]
            if len(parts) == 3:
                h,m,sec = parts
            elif len(parts) == 2:
                h,m,sec = 0, parts[0], parts[1]
            else:
                return np.nan
            return h*3600 + m*60 + sec
        except Exception:
            pass
    h=m=sec=0
    m1 = re.search(r"(\d+)\s*小?时", s)
    m2 = re.search(r"(\d+)\s*分", s)
    m3 = re.search(r"(\d+)\s*秒", s)
    if m1 or m2 or m3:
        if m1: h = int(m1.group(1))
        if m2: m = int(m2.group(1))
        if m3: sec = int(m3.group(1))
        return h*3600 + m*60 + sec
    return np.nan

def _enrich_comm_strict(df: pd.DataFrame, phone_map: Dict[str, Tuple[str, str]]) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame()
    d = clean_columns(df)
    if "对方号码" not in d.columns:
        return pd.DataFrame()

    if "对方姓名" not in d.columns: d["对方姓名"] = ""
    if "对方职务" not in d.columns: d["对方职务"] = ""

    raw_phone = d["对方号码"]
    uniq = pd.unique(raw_phone)
    norm_map = {v: normalize_phone_cell(v) for v in uniq}
    norm_phone = raw_phone.map(norm_map)

    name_dict = {k: v[0] for k, v in phone_map.items()}
    title_dict = {k: v[1] for k, v in phone_map.items()}

    mapped_name = norm_phone.map(name_dict).fillna("")
    mapped_title = norm_phone.map(title_dict).fillna("")
    d["对方姓名"] = np.where(mapped_name != "", mapped_name, d["对方姓名"].map(safe_str))
    d["对方职务"] = np.where(mapped_title != "", mapped_title, d["对方职务"].map(safe_str))

    ts = _compose_datetime_from_cols_relaxed(d)
    d["__ts__"] = ts
    if ts.notna().any():
        d["节日"] = _festival_series(ts)
        d["是否深夜(23–5)"] = _flag_late_night(ts).map({True: "是", False: ""})
    else:
        d["节日"] = ""
        d["是否深夜(23–5)"] = ""
    return d

def _mode_nonempty(series: pd.Series) -> str:
    s = series.fillna("").map(safe_str).str.strip()
    s = s[s != ""]
    if s.empty:
        return ""
    return s.value_counts().idxmax()

def _stats_by_phone(enriched_df: pd.DataFrame) -> pd.DataFrame:
    if enriched_df is None or enriched_df.empty:
        return pd.DataFrame()
    d = clean_columns(enriched_df)
    if "对方号码" not in d.columns:
        return pd.DataFrame()

    uniq = pd.unique(d["对方号码"])
    d["__对方号码__"] = d["对方号码"].map({v: normalize_phone_cell(v) for v in uniq})

    ts = d["__ts__"] if "__ts__" in d.columns else _compose_datetime_from_cols_relaxed(d)
    d["__ts__"] = ts

    nm = d["对方姓名"].map(safe_str) if "对方姓名" in d.columns else pd.Series([""]*len(d), index=d.index)
    title = d["对方职务"].map(safe_str) if "对方职务" in d.columns else pd.Series([""]*len(d), index=d.index)

    dur_col = next((c for c in ["通话时长","时长"] if c in d.columns), None)
    if dur_col:
        dur = d[dur_col].astype(str).str.strip()
        fast = dur.str.fullmatch(r"\d+(\.\d+)?([eE][+-]?\d+)?")
        dur_sec = pd.Series(np.nan, index=d.index, dtype=float)
        if fast.any():
            dur_sec.loc[fast] = dur.loc[fast].astype(float).values
        left = ~fast
        if left.any():
            dur_sec.loc[left] = dur.loc[left].apply(_parse_duration_to_seconds)
    else:
        dur_sec = pd.Series(np.nan, index=d.index, dtype=float)

    d["__dur_sec__"] = pd.to_numeric(dur_sec, errors="coerce")
    d["__姓名__"] = nm
    d["__职务__"] = title
    d["__非工作时间__"] = _flag_offwork(ts).astype(int)
    d["__深夜__"] = _flag_late_night(ts).astype(int)
    d["__通话3分钟以上__"] = (d["__dur_sec__"] >= 180).astype(int)
    fest = _festival_series(ts)

    # 统计主表：避免多次 groupby.apply，降低大通信表的运行时间与 pandas 兼容性风险。
    grp = d.groupby("__对方号码__", dropna=False)
    base = grp.size().rename("通信次数").to_frame()
    base["非工作时间通信次数"] = grp["__非工作时间__"].sum().astype(int)
    base["深夜通信次数(23–5)"] = grp["__深夜__"].sum().astype(int)
    base["通话≥3分钟次数"] = grp["__通话3分钟以上__"].sum().astype(int)
    base["姓名"] = grp["__姓名__"].agg(_mode_nonempty)
    base["职务"] = grp["__职务__"].agg(_mode_nonempty)
    base = base.reset_index().rename(columns={"__对方号码__": "对方号码"})

    # 节日次数：透视表（向量化）
    fest_df = pd.DataFrame({"对方号码": d["__对方号码__"], "节日": fest})
    fest_df = fest_df[fest_df["节日"].isin(FESTIVAL_NAMES)]
    if not fest_df.empty:
        pv = fest_df.assign(cnt=1).pivot_table(index="对方号码", columns="节日", values="cnt", aggfunc="sum", fill_value=0)
        pv.columns = [f"{c}通信次数" for c in pv.columns]
        pv = pv.reset_index()
        out = base.merge(pv, on="对方号码", how="left")
    else:
        out = base

    for fname in FESTIVAL_NAMES:
        coln = f"{fname}通信次数"
        if coln not in out.columns:
            out[coln] = 0
        out[coln] = pd.to_numeric(out[coln], errors="coerce").fillna(0).astype(int)

    out = out.sort_values(["通信次数","通话≥3分钟次数"], ascending=[False, False], kind="mergesort").reset_index(drop=True)
    return out

def load_and_enrich_communications_strict(root: Path, phone_to_name_title: Dict[str, Tuple[str, str]]) -> Dict[str, str]:
    STATE.COMM_STATS_BY_PERSON = {}
    STATE.COMM_RECORDS_BY_PERSON = {}
    STATE.COMM_STATS_ALL = None
    STATE.COMM_RECORDS_COUNT = 0

    if not phone_to_name_title:
        print("ℹ️ 未能从通讯录生成号码映射，跳过通信标注。")
        return {}

    files = [
        p for p in root.rglob("*.xls*")
        if (not p.name.startswith("~$")) and ("已标注" not in p.stem) and (not STATE.in_out_dir(p)) and COMM_FILE_RE.match(p.stem)
    ]
    STATE.COMM_FILES_COUNT = len(files)
    if not files:
        print(f"ℹ️ 未发现符合命名“*-{COMM_SUFFIX}.xlsx/.xls”的通信文件。")
        return {}

    name_to_title_out: Dict[str, str] = {}
    all_enriched_frames: List[pd.DataFrame] = []

    for p in files:
        print(f"📞 通信匹配：{p.name} ...")

        try:
            xls = pd.ExcelFile(p)
        except Exception as e:
            print("❌ 通信文件载入失败", p.name, e)
            continue

        frames: List[pd.DataFrame] = []
        name_map_file: Dict[str, str] = {}

        for sht in xls.sheet_names:
            try:
                hdr = _find_header_row_exact(xls, sht, ["对方号码"], 50)
                if hdr is None:
                    print(f"  • 跳过 {p.name}/{sht}：未找到表头（至少需要‘对方号码’）")
                    continue
                df0 = clean_columns(xls.parse(sheet_name=sht, header=hdr))
            except Exception as e:
                print("❌ 通信解析失败", f"{p.name}->{sht}", e)
                continue

            enriched = _enrich_comm_strict(df0, phone_to_name_title)
            if enriched.empty:
                continue
            
            # ================= 新增：从“本方姓名”列动态提取查询对象 =================
            if "本方姓名" in enriched.columns:
                query_objs = enriched["本方姓名"].map(safe_str).str.strip().replace("", "未知")
            else:
                fallback_name = safe_str(_person_from_comm_filename(p)).strip() or "未知"
                query_objs = pd.Series([fallback_name] * len(enriched), index=enriched.index)
            
            enriched.insert(0, "查询对象", query_objs)
            # ========================================================================

            if "__来源sheet__" not in enriched.columns:
                enriched.insert(0, "__来源sheet__", sht)

            frames.append(enriched)

            tmp = enriched[["对方姓名","对方职务"]].copy()
            tmp = tmp[(tmp["对方姓名"] != "") & (tmp["对方职务"] != "")]
            for nm, sub in tmp.groupby("对方姓名"):
                uniq_titles = list(dict.fromkeys(sub["对方职务"].map(safe_str).tolist()))
                name_map_file[nm] = "、".join(x for x in uniq_titles if x)

        if frames:
            merged = pd.concat(frames, ignore_index=True).replace({np.nan: ""})
            merged = merged.drop(columns=["__ts__"], errors="ignore")

            # 按“本方姓名”(即查询对象) 进行分组归并和输出标注文件，即使多个人在同一个文件中也能正确分发
            for person_guess, sub_df in merged.groupby("查询对象"):
                person_guess = str(person_guess)
                
                STATE.COMM_RECORDS_COUNT += len(sub_df)
                STATE.COMM_RECORDS_BY_PERSON[person_guess] = STATE.COMM_RECORDS_BY_PERSON.get(person_guess, 0) + len(sub_df)

                comm_mark_dir = _ensure_dir(out_comm_dir(person_guess) / "已标注")
                
                # 为防止重名文件互相覆盖，在文件名前拼接上对应的人名
                out_name = p.stem if person_guess in p.stem else f"{person_guess}-{p.stem}"
                save_df_auto_width(sub_df, comm_mark_dir / f"{out_name}-已标注", index=False, engine="openpyxl")
                print(f"✅ 通信标注导出：{comm_mark_dir / (out_name + '-已标注.xlsx')}")

            all_enriched_frames.append(merged)

        for k, v in name_map_file.items():
            if k in name_to_title_out and name_to_title_out[k]:
                exist = name_to_title_out[k].split("、")
                add = [x for x in v.split("、") if x not in exist]
                name_to_title_out[k] = "、".join(exist + add)
            else:
                name_to_title_out[k] = v

    if all_enriched_frames:
        merged_all = pd.concat(all_enriched_frames, ignore_index=True)
        
        # 统一生成每位人员汇总后的通信统计表
        for person_guess, sub_df in merged_all.groupby("查询对象"):
            person_guess = str(person_guess)
            stat_df = _stats_by_phone(sub_df)
            if stat_df is not None and not stat_df.empty:
                STATE.COMM_STATS_BY_PERSON[person_guess] = stat_df
                comm_stat_dir = _ensure_dir(out_comm_dir(person_guess) / "统计")
                save_df_auto_width(stat_df, comm_stat_dir / f"{person_guess}-汇总通信统计-按号码", index=False, engine="openpyxl")
                print(f"✅ 个人通信统计汇总导出：{comm_stat_dir / (person_guess + '-汇总通信统计-按号码.xlsx')}")

        # 全局统计生成
        stat_all = _stats_by_phone(merged_all)
        STATE.COMM_STATS_ALL = stat_all if (stat_all is not None and not stat_all.empty) else None
        
        comm_sum_stat_dir = _ensure_dir(out_comm_summary_dir() / "统计")
        
        if STATE.COMM_STATS_ALL is not None and not STATE.COMM_STATS_ALL.empty:
            save_df_auto_width(STATE.COMM_STATS_ALL, comm_sum_stat_dir / "ALL-通信统计-按号码", index=False, engine="openpyxl")
            print("✅ ALL通信统计汇总导出：分析结果/通信汇总/统计/ALL-通信统计-按号码.xlsx")

        # ================= 共同联系人分析 =================
        if "查询对象" in merged_all.columns and "对方号码" in merged_all.columns:
            print("🔍 正在跨文件分析共同联系人...")
            
            uniq_phones = pd.unique(merged_all["对方号码"])
            merged_all["__规范号码__"] = merged_all["对方号码"].map({v: normalize_phone_cell(v) for v in uniq_phones})
            
            grp = merged_all.groupby("__规范号码__", dropna=False)
            common_data = []
            
            for phone, g in grp:
                if not str(phone).strip():
                    continue
                targets = g["查询对象"].dropna().astype(str).str.strip()
                targets = targets[targets != ""]
                uniq_targets = sorted(targets.unique())
                
                if len(uniq_targets) >= 2:
                    nm = _mode_nonempty(g["对方姓名"]) if "对方姓名" in g.columns else ""
                    title = _mode_nonempty(g["对方职务"]) if "对方职务" in g.columns else ""
                    
                    row_data = {
                        "共同联系人号码": phone,
                        "姓名 (根据通讯录)": nm,
                        "职务 (根据通讯录)": title,
                        "关联核心对象数": len(uniq_targets),
                        "关联的核心对象": "、".join(uniq_targets),
                        "与所有对象总通信数": len(g)
                    }
                    
                    target_counts = targets.value_counts()
                    for target_name, count in target_counts.items():
                        row_data[f"与[{target_name}]通信次数"] = int(count)
                        
                    common_data.append(row_data)
            
            if common_data:
                df_common = pd.DataFrame(common_data)
                
                dynamic_cols = [c for c in df_common.columns if c.startswith("与[") and c.endswith("]通信次数")]
                df_common[dynamic_cols] = df_common[dynamic_cols].fillna(0).astype(int)
                
                base_cols = ["共同联系人号码", "姓名 (根据通讯录)", "职务 (根据通讯录)", "关联核心对象数", "关联的核心对象", "与所有对象总通信数"]
                final_cols = base_cols + sorted(dynamic_cols)
                df_common = df_common[final_cols]
                
                df_common.sort_values(["关联核心对象数", "与所有对象总通信数"], ascending=[False, False], inplace=True)
                
                out_path = comm_sum_stat_dir / "共同联系人分析"
                save_df_auto_width(df_common, out_path, index=False, engine="openpyxl")
                print(f"✅ 共同联系人分析导出：分析结果/通信汇总/统计/共同联系人分析.xlsx (共发现 {len(df_common)} 个)")
            else:
                print("ℹ️ 未发现共同联系人（未检测到同一个号码与多名调查对象均有通信）。")
        # =======================================================

    print(f"✅ 通信姓名映射生成 {len(name_to_title_out)} 条。")
    return name_to_title_out

# ========= 不动产（按文件名归属） =========
def _read_realestate_for_person_by_filename(root: Path, person: str) -> Tuple[List[Dict[str, Any]], int]:
    person = safe_str(person).strip()
    if not person or person == "未知":
        return [], 0

    files = []
    for p in root.rglob("*.xlsx"):
        if p.name.startswith("~$") or STATE.in_out_dir(p):
            continue
        if p.stem.startswith(f"{person}-全省不动产"):
            files.append(p)

    records: List[Dict[str, Any]] = []
    for p in files:
        try:
            xls = pd.ExcelFile(p)
        except Exception:
            continue
        for sht in xls.sheet_names:
            try:
                df = xls.parse(sht).dropna(how="all")
            except Exception:
                continue
            df = clean_columns(df)
            for _, row in df.iterrows():
                loc = safe_str(row.get("房屋坐落", "")).strip()
                area = safe_str(row.get("建筑面积", "")).strip()
                price = safe_str(row.get("交易价格（万元）", "")).strip()
                regt = safe_str(row.get("登记时间", "")).strip()
                if not (loc or area or price or regt):
                    continue
                records.append({"来源文件": p.name, "sheet": sht, "房屋坐落": loc, "建筑面积": area, "交易价格（万元）": price, "登记时间": regt})
    return records, len(files)


# ========= 分析研判交易流水模板识别/转换 =========
ANALYSIS_JUDGMENT_TXN_REQUIRED = ["本方名称", "对方名称", "交易时间", "交易金额", "借贷类型"]
ANALYSIS_JUDGMENT_TXN_HEADERS = [
    "所属群组", "本方名称", "本方展示卡号", "本方所属银行", "本方开户地",
    "对方名称", "对方展示卡号", "对方所属银行", "对方开户地", "交易时间",
    "星期", "交易金额", "币种", "账户余额", "借贷类型", "对方工作单位",
    "对方职务", "摘要", "备注/注释", "交易类型", "交易机构名称", "交易网点",
    "柜员编号", "MAC地址", "IP地址",
]


def _format_txn_time_any(v: Any) -> str:
    s = safe_str(v).strip()
    r = _parse_compact_datetime(s)
    if r:
        return r
    tt = pd.to_datetime(s, errors="coerce")
    return tt.strftime("%Y-%m-%d %H:%M:%S") if pd.notna(tt) else (s or "wrong")


def _normalize_debit_credit_flag(x: Any) -> str:
    """把不同来源的借贷方向统一为“进/出”。"""
    s = safe_str(x).strip()
    if not s:
        return ""
    s_up = s.upper()
    if s_up in {"1", "D", "DR", "DEBIT", "借", "出", "支出", "转出", "付款", "付", "消费", "取现"}:
        return "出"
    if s_up in {"2", "C", "CR", "CREDIT", "贷", "进", "收入", "转入", "收款", "收", "入账", "存入"}:
        return "进"
    if re.search(r"支出|转出|付款|消费|取现|借方|付出", s):
        return "出"
    if re.search(r"收入|转入|收款|入账|存入|贷方|收到", s):
        return "进"
    return s


def _parse_money_cell(x: Any, default: float = 0.0, abs_value: bool = False) -> float:
    """安全解析金额单元格。

    兼容“分析研判-交易流水”常见文本金额：
    7,000.00、7，000.00、￥7,000.00、人民币7,000元、(7,000.00)、-7,000.00 等。
    解析失败时返回 default，避免 NaN 进入 int()/分组/评分逻辑。
    """
    try:
        if x is None:
            return default
        if isinstance(x, (int, float, np.integer, np.floating)):
            val = float(x)
            if not np.isfinite(val):
                return default
            return abs(val) if abs_value else val
    except Exception:
        pass

    s = safe_str(x).strip()
    if not s:
        return default

    s = unicodedata.normalize("NFKC", s)
    s = s.replace("，", ",").replace(" ", "").replace("\u00A0", "")
    lowered = s.lower()
    if lowered in {"nan", "none", "null", "--", "-", "无", "空"}:
        return default

    negative = False
    if re.fullmatch(r"\([^()]+\)", s):
        negative = True
        s = s[1:-1]

    # 去掉常见币种/单位/方向文字，只保留可解析的数字结构。
    s = re.sub(r"人民币|RMB|CNY|￥|¥|元|圆|整|金额|交易金额|收入|支出|转入|转出|借方|贷方|借|贷", "", s, flags=re.IGNORECASE)
    s = s.replace(",", "")

    m = re.search(r"[-+]?\d+(?:\.\d+)?", s)
    if not m:
        return default
    try:
        val = float(m.group(0))
        if negative and val > 0:
            val = -val
        if not np.isfinite(val):
            return default
        return abs(val) if abs_value else val
    except Exception:
        return default


def _safe_money_series(s: Any, default: float = 0.0, abs_value: bool = False) -> pd.Series:
    """金额列安全转换为 float Series，自动处理文本金额和 NaN。"""
    if isinstance(s, pd.Series):
        return s.map(lambda v: _parse_money_cell(v, default=default, abs_value=abs_value)).astype(float)
    return pd.Series([_parse_money_cell(s, default=default, abs_value=abs_value)], dtype=float)

def _is_analysis_judgment_source_row(row: Any) -> bool:
    """识别单行是否来自“分析研判-交易流水”模板。"""
    vals = []
    for col in ["查询项", "数据来源类型", "来源文件", "来源sheet", "所属群组"]:
        try:
            vals.append(safe_str(row.get(col, "")))
        except Exception:
            pass
    joined = " ".join(vals)
    return "分析研判" in joined


def _mark_analysis_judgment_source(df: pd.DataFrame) -> pd.Series:
    """返回布尔序列：True 表示该行来源为“分析研判-交易流水”。"""
    if df is None or df.empty:
        return pd.Series([], dtype=bool)
    parts = []
    for col in ["查询项", "数据来源类型", "来源文件", "来源sheet", "所属群组"]:
        if col in df.columns:
            parts.append(df[col].map(safe_str))
    if not parts:
        return pd.Series(False, index=df.index)
    joined = parts[0].copy()
    for s in parts[1:]:
        joined = joined + " " + s
    return joined.str.contains("分析研判", na=False)


def _txn_second_key_series(s: pd.Series) -> pd.Series:
    """交易时间去重键：能解析的统一到秒，不能解析的保留原始文本。"""
    raw = s.map(_format_txn_time_any) if s is not None else pd.Series([], dtype=object)
    dt = pd.to_datetime(raw, errors="coerce")
    key = raw.map(safe_str).str.strip()
    ok = dt.notna()
    if ok.any():
        key.loc[ok] = dt.loc[ok].dt.floor("s").dt.strftime("%Y-%m-%d %H:%M:%S")
    return key.fillna("")


def _looks_like_analysis_judgment_txn(df: pd.DataFrame) -> bool:
    # 只看表头，不因样例表暂时没有数据行而漏识别。
    if df is None:
        return False
    cols = {safe_str(c).strip() for c in df.columns}
    return set(ANALYSIS_JUDGMENT_TXN_REQUIRED).issubset(cols)


def _infer_person_from_analysis_judgment_filename(p: Path) -> str:
    stem = safe_str(p.stem).strip()
    for sep in ["-分析研判", "_分析研判", "-交易流水", "_交易流水"]:
        if sep in stem:
            name = stem.split(sep, 1)[0].strip()
            if name:
                return name
    if "-" in stem:
        return stem.split("-", 1)[0].strip()
    return stem or "未知"


def _ensure_bank_template_columns(df: pd.DataFrame) -> pd.DataFrame:
    """补齐银行模板基础列，同时保留新模板带来的额外字段。"""
    d = clean_columns(df)
    for col in TEMPLATE_COLS:
        if col not in d.columns:
            d[col] = ""
    base_cols = [c for c in TEMPLATE_COLS if c in d.columns]
    extra_cols = [c for c in d.columns if c not in base_cols]
    return d[base_cols + extra_cols]


def _convert_analysis_judgment_txn_df(df: pd.DataFrame, p: Path, sheet_name: str) -> pd.DataFrame:
    """将“分析研判-交易流水”表转换为既有银行模板字段。"""
    src = clean_columns(df).dropna(how="all").copy()
    if src.empty:
        return pd.DataFrame()

    # 删除完全没有交易要素的空白行，避免样例表只有表头时进入后续分析。
    key_cols = [c for c in ["本方名称", "对方名称", "交易时间", "交易金额", "借贷类型"] if c in src.columns]
    if key_cols:
        nonempty = src[key_cols].apply(lambda col: col.map(safe_str).str.strip() != "").any(axis=1)
        src = src[nonempty].copy()
    if src.empty:
        return pd.DataFrame()

    out = pd.DataFrame(index=src.index)
    fallback_person = _infer_person_from_analysis_judgment_filename(p)

    out["查询对象"] = src.get("本方名称", pd.Series([fallback_person] * len(src), index=src.index)).map(safe_str).str.strip()
    out["查询对象"] = out["查询对象"].replace({"": fallback_person, "未知": fallback_person, "nan": fallback_person, "NaN": fallback_person})
    out["反馈单位"] = src.get("本方所属银行", "")
    out["查询项"] = "分析研判交易流水"
    out["数据来源类型"] = "分析研判"
    out["查询账户"] = src.get("本方展示卡号", "")
    out["查询卡号"] = src.get("本方展示卡号", "")
    out["交易类型"] = src.get("交易类型", "")
    out["借贷标志"] = src.get("借贷类型", "").map(_normalize_debit_credit_flag) if "借贷类型" in src.columns else ""
    out["币种"] = src.get("币种", "")
    out["交易金额"] = _safe_money_series(src.get("交易金额", pd.Series([0] * len(src), index=src.index)), default=0.0, abs_value=True).round(2)
    out["账户余额"] = _safe_money_series(src.get("账户余额", pd.Series([np.nan] * len(src), index=src.index)), default=np.nan, abs_value=False).round(2)
    out["交易时间"] = src.get("交易时间", "").map(_format_txn_time_any) if "交易时间" in src.columns else "wrong"
    out["交易流水号"] = [f"{p.stem}-{sheet_name}-{i+1}" for i in range(len(src))]
    out["本方账号"] = src.get("本方展示卡号", "")
    out["本方卡号"] = src.get("本方展示卡号", "")
    out["交易对方姓名"] = src.get("对方名称", "")
    out["交易对方账户"] = src.get("对方展示卡号", "")
    out["交易对方卡号"] = src.get("对方展示卡号", "")
    out["交易对方证件号码"] = ""
    out["交易对手余额"] = ""
    out["交易对方账号开户行"] = src.get("对方所属银行", "")
    out["交易摘要"] = src.get("摘要", "")
    out["交易网点名称"] = src.get("交易机构名称", src.get("交易网点", ""))
    out["交易网点代码"] = src.get("交易网点", "")
    out["日志号"] = ""
    out["传票号"] = ""
    out["凭证种类"] = ""
    out["凭证号"] = ""
    out["现金标志"] = ""
    out["终端号"] = ""
    out["交易是否成功"] = ""
    out["交易发生地"] = src.get("本方开户地", "")
    out["商户名称"] = src.get("对方工作单位", "")
    out["商户号"] = ""
    out["IP地址"] = src.get("IP地址", "")
    out["MAC"] = src.get("MAC地址", "")
    out["交易柜员号"] = src.get("柜员编号", "")
    out["备注"] = src.get("备注/注释", "")

    # 保留“分析研判”表的原始特征列，供三类线索文本检索和人工复核使用。
    out["所属群组"] = src.get("所属群组", "")
    out["本方开户地"] = src.get("本方开户地", "")
    out["对方开户地"] = src.get("对方开户地", "")
    out["对方工作单位"] = src.get("对方工作单位", "")
    out["对方职务"] = src.get("对方职务", "")
    out["摘要"] = src.get("摘要", "")
    out["备注/注释"] = src.get("备注/注释", "")
    out["交易机构名称"] = src.get("交易机构名称", "")
    out["来源文件"] = p.name
    out["来源sheet"] = sheet_name

    return _ensure_bank_template_columns(out).replace({np.nan: ""})


def _read_analysis_judgment_txn_file(p: Path) -> Tuple[pd.DataFrame, bool]:
    """读取“分析研判-交易流水”新模板。返回：(转换结果, 是否识别为该模板)。"""
    try:
        xls = pd.ExcelFile(p)
    except Exception:
        return pd.DataFrame(), False

    frames: List[pd.DataFrame] = []
    matched = False
    for sht in xls.sheet_names:
        try:
            hdr = _find_header_row_contains(xls, sht, ANALYSIS_JUDGMENT_TXN_REQUIRED, scan_rows=60)
            if hdr is None:
                continue
            df0 = clean_columns(xls.parse(sheet_name=sht, header=hdr))
            if not _looks_like_analysis_judgment_txn(df0):
                continue
            matched = True
            conv = _convert_analysis_judgment_txn_df(df0, p, sht)
            if conv is not None and not conv.empty:
                frames.append(conv)
        except Exception as e:
            print("❌ 分析研判交易流水解析失败", f"{p.name}->{sht}", e)

    if frames:
        return pd.concat(frames, ignore_index=True), True
    return pd.DataFrame(), matched


# ========= 交易流水合并 =========
def merge_all_txn(root_dir: str) -> pd.DataFrame:
    root = Path(root_dir).expanduser().resolve()

    comm_files = [
        p for p in root.rglob("*.xls*")
        if (not p.name.startswith("~$")) and ("已标注" not in p.stem) and (not STATE.in_out_dir(p)) and COMM_FILE_RE.match(p.stem)
    ]
    if comm_files:
        print(f"📁 检测到 {len(comm_files)} 个“-通信”文件，将加载通讯录并进行通信标注。")
        STATE.CONTACT_PHONE_TO_NAME_TITLE = load_contacts_phone_map_strict(root)
        STATE.CALLLOG_NAME_TO_TITLE = load_and_enrich_communications_strict(root, STATE.CONTACT_PHONE_TO_NAME_TITLE)
        annotate_hotel_and_flight_files(root, STATE.CALLLOG_NAME_TO_TITLE)
    else:
        STATE.CONTACT_PHONE_TO_NAME_TITLE = {}
        STATE.CALLLOG_NAME_TO_TITLE = {}
        print(f"ℹ️ 未发现符合命名“*-{COMM_SUFFIX}.xlsx/.xls”的通信文件，本次不读取通讯录，也不做通信标注；同时不做“同住/同机”职务匹配。")

    # 银行交易流水入口：保留原“网上银行”模板，同时兼容“分析研判-交易流水”模板。
    china_files = [
        p for p in root.rglob("*-*-交易流水.xls*")
        if (not p.name.startswith("~$")) and (not STATE.in_out_dir(p))
    ]

    print(f"✅ 银行交易流水文件 {len(china_files)} 个；通信映射 {len(STATE.CALLLOG_NAME_TO_TITLE)} 条。")

    dfs: List[pd.DataFrame] = []
    processed_files: set[Path] = set()

    def _append_and_mark(df: pd.DataFrame, p: Path):
        if df is not None and not df.empty:
            dfs.append(df); processed_files.add(p)

    # 网银模板直接读；“分析研判-交易流水”模板先转换为统一银行模板再进入同一分析流程。
    for p in china_files:
        if p in processed_files:
            continue
        print(f"正在处理 {p.name} ...")
        try:
            df_analysis, is_analysis_judgment = _read_analysis_judgment_txn_file(p)
            if is_analysis_judgment:
                if df_analysis is not None and not df_analysis.empty:
                    print(f"  • 识别为“分析研判-交易流水”模板，转换后 {len(df_analysis)} 条。")
                    _append_and_mark(df_analysis, p)
                else:
                    print("  • 识别为“分析研判-交易流水”模板，但未发现有效交易数据。")
                continue

            df = pd.read_excel(p, dtype={"查询卡号": str, "查询账户": str, "交易对方证件号码": str, "本方账号": str, "本方卡号": str})
            df = _ensure_bank_template_columns(df)
            if "交易时间" in df.columns:
                df["交易时间"] = df["交易时间"].map(_format_txn_time_any)
            df["来源文件"] = p.name
            df["数据来源类型"] = "网上银行"
            _append_and_mark(df, p)
        except Exception as e:
            print("❌", p.name, e)

    # 已按要求移除农商、泰隆、民泰、农行线下、建行线下及 CSV 的交易处理流程。

    print("文件读取完成，正在整合……")
    if not dfs:
        STATE.RAW_TXN_COUNT = STATE.DUP_TXN_COUNT = STATE.FINAL_TXN_COUNT = 0
        return pd.DataFrame(columns=TEMPLATE_COLS)

    raw_txn = pd.concat(dfs, ignore_index=True)
    STATE.RAW_TXN_COUNT = len(raw_txn)

    # 金额列可能存在空白、"NaN"、inf 等异常值；统一转为 0，避免后续 int()/分组评分报错。
    raw_txn["交易金额"] = _safe_money_series(raw_txn.get("交易金额", pd.Series([0] * len(raw_txn), index=raw_txn.index)), default=0.0, abs_value=True).round(2)
    if "数据来源类型" not in raw_txn.columns:
        raw_txn["数据来源类型"] = ""
    raw_txn["数据来源类型"] = raw_txn["数据来源类型"].map(safe_str).str.strip()
    raw_txn.loc[_mark_analysis_judgment_source(raw_txn), "数据来源类型"] = "分析研判"
    raw_txn.loc[raw_txn["数据来源类型"] == "", "数据来源类型"] = "网上银行"

    # 统一借贷方向和交易时间后，再做跨模板去重。
    # 规则：交易时间精确到秒 + 交易金额 + 借贷标志 相同即视为重复；
    #       优先保留非“分析研判-交易流水”来源，删除“分析研判”来源的重复记录。
    raw_txn["借贷标志"] = raw_txn.get("借贷标志", "").map(_normalize_debit_credit_flag)
    raw_txn["交易时间"] = raw_txn.get("交易时间", "").map(_format_txn_time_any)
    raw_txn["__去重时间__"] = _txn_second_key_series(raw_txn["交易时间"])
    raw_txn["__去重金额__"] = _safe_money_series(raw_txn["交易金额"], default=0.0, abs_value=True).round(2)
    raw_txn["__去重借贷标志__"] = raw_txn["借贷标志"].map(_normalize_debit_credit_flag).map(safe_str).str.strip()
    raw_txn["__是否分析研判__"] = _mark_analysis_judgment_source(raw_txn)
    raw_txn["__保留优先级__"] = raw_txn["__是否分析研判__"].map({False: 0, True: 1}).fillna(1).astype(int)
    raw_txn["__原始顺序__"] = np.arange(len(raw_txn))

    dedup_keys = ["__去重时间__", "__去重金额__", "__去重借贷标志__"]
    ordered = raw_txn.sort_values(["__保留优先级__", "__原始顺序__"], ascending=[True, True], kind="mergesort")
    dup_mask_ordered = ordered.duplicated(subset=dedup_keys, keep="first")
    dup_df = ordered[dup_mask_ordered].copy()
    if not dup_df.empty:
        dup_df["重复判断依据"] = (
            "交易时间精确到秒=" + dup_df["__去重时间__"].map(safe_str)
            + "；交易金额=" + dup_df["__去重金额__"].map(lambda x: f"{float(x):.2f}" if pd.notna(x) else "0.00")
            + "；借贷标志=" + dup_df["__去重借贷标志__"].map(safe_str)
        )
        dup_df["去重保留规则"] = np.where(dup_df["__是否分析研判__"], "与非分析研判数据重复，删除分析研判记录", "同键重复，按原始读取顺序删除后出现记录")

    STATE.DUP_TXN_COUNT = int(dup_mask_ordered.sum())
    if not dup_df.empty:
        export_dup = dup_df.drop(columns=[c for c in dup_df.columns if c.startswith("__")], errors="ignore")
        save_df_auto_width(export_dup, "所有人-重复交易流水", index=False, engine="openpyxl")
        print(f"✅ 已导出重复交易流水：{len(export_dup)} 条 -> 分析结果/所有人-重复交易流水.xlsx")

    before = len(raw_txn)
    all_txn = ordered[~dup_mask_ordered].sort_values("__原始顺序__", kind="mergesort").reset_index(drop=True)
    all_txn = all_txn.drop(columns=[c for c in all_txn.columns if c.startswith("__")], errors="ignore")
    if "数据来源类型" not in all_txn.columns:
        all_txn["数据来源类型"] = ""
    all_txn["数据来源类型"] = all_txn["数据来源类型"].map(safe_str).str.strip()
    all_txn.loc[_mark_analysis_judgment_source(all_txn), "数据来源类型"] = "分析研判"
    all_txn.loc[all_txn["数据来源类型"] == "", "数据来源类型"] = "网上银行"
    STATE.FINAL_TXN_COUNT = len(all_txn)
    removed = before - len(all_txn)
    if removed:
        print(f"🧹 跨模板去重 {removed} 条（规则：交易时间精确到秒 + 交易金额 + 借贷标志；优先保留非分析研判数据）.")

    ts = pd.to_datetime(all_txn["交易时间"], errors="coerce")
    all_txn.insert(0, "__ts__", ts)
    all_txn.sort_values("__ts__", inplace=True, kind="mergesort")
    all_txn["序号"] = range(1, len(all_txn) + 1)
    all_txn.drop(columns="__ts__", inplace=True)

    all_txn["借贷标志"] = all_txn["借贷标志"].apply(_normalize_debit_credit_flag)

    bins = [-np.inf, 2000, 5000, 20000, 50000, np.inf]
    labels = ["2000以下","2000-5000","5000-20000","20000-50000","50000以上"]
    all_txn["金额区间"] = pd.cut(_safe_money_series(all_txn["交易金额"], default=0.0, abs_value=True), bins=bins, labels=labels, right=False, include_lowest=True)

    weekday_map = {0:"星期一",1:"星期二",2:"星期三",3:"星期四",4:"星期五",5:"星期六",6:"星期日"}
    wk = pd.Series(index=all_txn.index, dtype=object)
    mask = ts.notna()
    wk.loc[mask] = ts.dt.weekday.map(weekday_map)
    wk.loc[~mask] = "wrong"
    all_txn["星期"] = wk

    dates = ts.dt.date
    status = pd.Series(index=all_txn.index, dtype=object)
    unique_dates = pd.unique(dates[mask])

    @lru_cache(maxsize=None)
    def _day_status(d) -> str:
        try:
            return "节假日" if is_holiday(d) else ("工作日" if is_workday(d) else "周末")
        except Exception:
            dd = _dt.datetime.combine(d, _dt.time())
            return "周末" if dd.weekday() >= 5 else "工作日"

    if len(unique_dates):
        mapd = {d: _day_status(d) for d in unique_dates}
        status.loc[mask] = dates.loc[mask].map(mapd)
    status.loc[~mask] = "wrong"
    all_txn["节假日"] = status

    mapped_titles = all_txn["交易对方姓名"].map(STATE.CALLLOG_NAME_TO_TITLE).fillna("")
    if "对方职务" in all_txn.columns:
        existed_titles = all_txn["对方职务"].map(safe_str).str.strip()
        all_txn["对方职务"] = np.where(existed_titles != "", existed_titles, mapped_titles)
    else:
        all_txn["对方职务"] = mapped_titles
    cols = list(all_txn.columns)
    if "交易对方姓名" in cols and "对方职务" in cols:
        cols.remove("对方职务")
        insert_at = cols.index("交易对方姓名") + 1
        cols = cols[:insert_at] + ["对方职务"] + cols[insert_at:]
        all_txn = all_txn[cols]

    save_df_auto_width(all_txn, "所有人-合并交易流水", index=False, engine="openpyxl")
    print("✅ 已导出：分析结果/所有人-合并交易流水.xlsx")
    return all_txn


# ========= 现金交易识别 =========
CASH_TEXT_KEYWORDS_RE = re.compile(r"现存|现取|卡存|卡取|ATM", re.IGNORECASE)

def _series_text(df: pd.DataFrame, col: str) -> pd.Series:
    if df is not None and col in df.columns:
        return df[col].map(safe_str)
    return pd.Series([""] * (len(df) if df is not None else 0), index=(df.index if df is not None else None), dtype=object)

def _cash_signal_mask(df: pd.DataFrame) -> pd.Series:
    """识别现金/存取现交易。

    原逻辑保留：现金标志含“现”或数值为 1、交易类型含“柜面/现”。
    新增逻辑：交易摘要、摘要、备注、注释列中出现“现存、现取、卡存、卡取、ATM”。
    """
    if df is None or df.empty:
        return pd.Series([], dtype=bool)

    idx = df.index
    cash_flag = _series_text(df, "现金标志").str.contains("现", na=False) | (pd.to_numeric(_series_text(df, "现金标志"), errors="coerce") == 1)
    txn_type = _series_text(df, "交易类型").str.contains("柜面|现", na=False)

    text_mask = pd.Series(False, index=idx)
    for col in ["交易摘要", "摘要", "备注", "注释","备注/注释"]:
        if col in df.columns:
            text_mask = text_mask | _series_text(df, col).str.contains(CASH_TEXT_KEYWORDS_RE, na=False)

    return cash_flag | txn_type | text_mask

# ========= 分析输出（个人银行） =========
def analysis_txn(df: pd.DataFrame) -> None:
    if df.empty:
        return
    df = df.copy()
    df["交易时间"] = pd.to_datetime(df["交易时间"], errors="coerce")
    df["交易金额"] = pd.to_numeric(df["交易金额"], errors="coerce").replace([np.inf, -np.inf], np.nan).fillna(0)
    person = safe_str(df["查询对象"].iat[0]).strip() or "未知"
    prefix = out_bank_dir(person)

    out_df = df[df["借贷标志"] == "出"]
    in_df = df[df["借贷标志"] == "进"]
    counts = df["金额区间"].value_counts()

    summary = pd.DataFrame([{
        "交易次数": len(df),
        "交易金额": df["交易金额"].sum(skipna=True),
        "流出额": out_df["交易金额"].sum(skipna=True),
        "流入额": in_df["交易金额"].sum(skipna=True),
        "单笔最大支出": out_df["交易金额"].max(skipna=True),
        "单笔最大收入": in_df["交易金额"].max(skipna=True),
        "净流入": in_df["交易金额"].sum(skipna=True) - out_df["交易金额"].sum(skipna=True),
        "最后交易时间": df["交易时间"].max(),
        "0-2千次数": counts.get("2000以下", 0),
        "2千-5千次数": counts.get("2000-5000", 0),
        "5千-2万次数": counts.get("5000-20000", 0),
        "2万-5万次数": counts.get("20000-50000", 0),
        "5万以上次数": counts.get("50000以上", 0),
    }])
    save_df_auto_width(summary, prefix / f"0{person}-资产分析", index=False, engine="openpyxl")

    cash = df[_cash_signal_mask(df) & (pd.to_numeric(df["交易金额"], errors="coerce").abs() >= 10_000)]
    save_df_auto_width(cash, prefix / f"1{person}-存取现1万以上", index=False, engine="openpyxl")

    big = df[pd.to_numeric(df["交易金额"], errors="coerce") >= 500_000]
    save_df_auto_width(big, prefix / f"1{person}-大额资金50万以上", index=False, engine="openpyxl")

    src = df.copy()
    src["is_in"] = src["借贷标志"] == "进"
    src["signed_amt"] = pd.to_numeric(src["交易金额"], errors="coerce") * src["is_in"].map({True: 1, False: -1})
    src["in_amt"] = pd.to_numeric(src["交易金额"], errors="coerce").where(src["is_in"], 0)

    src = (src.groupby("交易对方姓名", dropna=False).agg(
        交易金额=("交易金额", "sum"),
        交易次数=("交易金额", "size"),
        流入额=("in_amt", "sum"),
        净流入=("signed_amt", "sum"),
        单笔最大收入=("in_amt", "max"),
    ).reset_index())

    total = src["流入额"].sum()
    src["流入比%"] = src["流入额"] / total * 100 if total else 0

    name_to_title = (df[["交易对方姓名","对方职务"]].dropna().drop_duplicates().set_index("交易对方姓名")["对方职务"].to_dict())
    src.insert(1, "对方职务", src["交易对方姓名"].map(name_to_title).fillna(""))

    src.sort_values("流入额", ascending=False, inplace=True)
    save_df_auto_width(src, prefix / f"1{person}-资金来源分析", index=False, engine="openpyxl")

def make_partner_summary(df: pd.DataFrame) -> None:
    if df.empty:
        return
    person = safe_str(df["查询对象"].iat[0]).strip() or "未知"
    prefix = out_bank_dir(person)

    d = df.copy()
    d["交易金额"] = pd.to_numeric(d["交易金额"], errors="coerce").replace([np.inf, -np.inf], np.nan).fillna(0)
    d["is_in"] = d["借贷标志"] == "进"
    d["abs_amt"] = d["交易金额"].abs()
    d["signed_amt"] = d["交易金额"] * d["is_in"].map({True: 1, False: -1})
    d["in_amt"] = d["交易金额"].where(d["is_in"], 0)
    d["out_amt"] = d["交易金额"].where(~d["is_in"], 0)
    d["gt10k"] = (d["abs_amt"] >= 10_000).astype(int)

    summ = (d.groupby(["查询对象","交易对方姓名"], dropna=False).agg(
        交易次数=("交易金额", "size"),
        交易金额=("abs_amt", "sum"),
        万元以上交易次数=("gt10k", "sum"),
        净收入=("signed_amt", "sum"),
        转入笔数=("is_in", "sum"),
        转入金额=("in_amt", "sum"),
        转出笔数=("is_in", lambda x: (~x).sum()),
        转出金额=("out_amt", "sum"),
    ).reset_index().rename(columns={"查询对象":"姓名","交易对方姓名":"对方姓名"}))

    name_to_title = (d[["交易对方姓名","对方职务"]].drop_duplicates().set_index("交易对方姓名")["对方职务"].to_dict())
    summ.insert(2, "对方职务", summ["对方姓名"].map(name_to_title).fillna(""))

    total = summ.groupby("姓名")["交易金额"].transform("sum")
    summ["交易占比%"] = np.where(total > 0, summ["交易金额"] / total * 100, 0)

    summ.sort_values(["姓名","交易金额"], ascending=[True, False], inplace=True)
    save_df_auto_width(summ, prefix / f"2{person}-交易对手分析", index=False, engine="openpyxl")

    comp = summ[summ["对方姓名"].map(safe_str).str.contains("公司", na=False)]
    save_df_auto_width(comp, prefix / f"3{person}-与公司相关交易频次分析", index=False, engine="openpyxl")



# ========= 三类线索识别（对象画像 + 分值研判） =========
CLUE_TYPES = ["疑似情感关系线索", "疑似资金代持/白手套线索", "疑似利益输送线索"]
CLUE_TEXT_COLS = ["交易摘要", "摘要", "备注", "注释", "备注/注释", "商户名称", "交易类型"]
COMPANY_KEYWORDS_RE = re.compile(r"公司|集团|厂|店|商行|经营部|个体|合作社|工程|建设|贸易|商贸|科技|咨询|服务|管理|投资|置业|房地产|劳务|建筑|装饰|材料|运输|物流", re.IGNORECASE)
EMOTION_KEYWORDS_RE = re.compile(
    r"红包|生日|礼物|礼金|520|521|1314|七夕|情人节|亲爱|宝贝|宝宝|老婆|老公|媳妇|对象|情人|爱人|恋人|爱你|想你|思念|纪念日|约会|玫瑰|鲜花|口红|香水|转给你花|给你买|么么|想见你|晚安|早安|亲亲|宝儿|乖|抱抱",
    re.IGNORECASE,
)
LIFECARE_KEYWORDS_RE = re.compile(
    r"生活费|房租|租金|水电|物业|买衣服|衣服|化妆品|护肤|口红|香水|医药费|看病|吃饭|饭钱|奶茶|咖啡|电影|旅游|酒店|住宿|宾馆|打车|车费|机票|高铁|花费|零花|给你用|给你花|补贴|家用",
    re.IGNORECASE,
)
HOLDING_KEYWORDS_RE = re.compile(
    r"代持|白手套|代收|代付|代转|周转|走账|过账|过桥|备用金|借款|还款|垫付|暂存|暂借|帮转|帮收|帮付|刷流水|通道|中转|倒账|套现|拆借|归还|挂账",
    re.IGNORECASE,
)
BENEFIT_KEYWORDS_RE = re.compile(
    r"好处费|感谢费|辛苦费|茶水费|回扣|返点|返利|介绍费|协调费|打点|关系费|咨询费|服务费|劳务费|居间|中介费|项目款|工程款|管理费|招待费|赞助费|活动费|顾问费|分红|提成|返点款|酬谢",
    re.IGNORECASE,
)
PROJECT_KEYWORDS_RE = re.compile(r"项目|工程|采购|招标|投标|中标|审批|许可|工程款|材料款|咨询|服务|劳务|管理|合同|货款|保证金|押金", re.IGNORECASE)
SPECIAL_EMOTION_AMOUNTS = {5.20, 5.21, 13.14, 52.0, 52.1, 66.66, 88.88, 99.99, 100.1, 520, 521, 666, 888, 999, 1314, 1314.52, 5200, 5201, 5210, 6666, 8888, 9999, 13140, 1314520}
SPECIAL_EMOTION_AMOUNT_TAILS = ("520", "521", "1314", "999", "888", "666")

CLUE_RULE_DEFINITIONS = [
    # 情感关系：单项分值降低，强调多规则、多证据叠加，避免单笔特殊金额直接推高。
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-01", "规则说明": "特殊表达金额：520、521、1314、999、888、666等固定金额或尾数表达。", "计分方式": "每笔2分，最高12分", "分值": 2, "上限": 12},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-02", "规则说明": "亲密/情感文本：摘要、备注、商户名称等含红包、生日、礼物、情人、亲昵称谓、纪念日、鲜花等。", "计分方式": "每笔4分，最高20分", "分值": 4, "上限": 20},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-03", "规则说明": "情感敏感日期：2月14日、5月20日、5月21日、七夕、生日/纪念日文本等日期附近资金往来。", "计分方式": "每笔3分，最高15分", "分值": 3, "上限": 15},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-04", "规则说明": "敏感时段：自然人对象在非工作时间、深夜、节假日多次资金往来。", "计分方式": "按非工作/深夜/节假日次数累计，最高14分", "分值": 1, "上限": 14},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-05", "规则说明": "资金叠加通信：同名对象存在高频通信、深夜通信、节日通信或较长通话。", "计分方式": "按通信强度累计，最高18分", "分值": 1, "上限": 18},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-06", "规则说明": "共同轨迹：同名对象出现在旅馆同住记录。", "计分方式": "每次18分，最高36分", "分值": 18, "上限": 36},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-07", "规则说明": "共同出行：同名对象出现在同机记录。", "计分方式": "每次8分，最高16分", "分值": 8, "上限": 16},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-08", "规则说明": "小额高频：自然人对象多笔小额往来，疑似日常性支持或亲密关系往来。", "计分方式": "从第4笔起每笔1分，最高12分", "分值": 1, "上限": 12},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-09", "规则说明": "双向小额互转：自然人对象既有小额转入又有小额转出。", "计分方式": "按双向小额笔数累计，最高12分", "分值": 2, "上限": 12},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-10", "规则说明": "长期持续往来：自然人对象跨多月仍有资金往来。", "计分方式": "从第3个月起累计，最高12分", "分值": 3, "上限": 12},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-11", "规则说明": "生活照顾类文本：摘要、备注含生活费、房租、衣服、化妆品、吃饭、旅游、住宿、打车等。", "计分方式": "每笔4分，最高16分", "分值": 4, "上限": 16},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-12", "规则说明": "单向供养/支持：向自然人对象转出为主，且多笔小额或中额支出。", "计分方式": "按转出笔数、金额和占比累计，最高14分", "分值": 2, "上限": 14},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-13", "规则说明": "自然人资金占比较高：同一自然人占本人资金往来比例较高。", "计分方式": "按占比累计，最高10分", "分值": 2, "上限": 10},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-14", "规则说明": "组合特征：特殊金额同时叠加亲密文本、敏感日期、非工作/深夜/节假日等。", "计分方式": "每类组合2分，最高12分", "分值": 2, "上限": 12},
    {"线索类型": "疑似情感关系线索", "规则编号": "EMO-15", "规则说明": "轨迹/通信/资金交叉：资金往来对象同时出现通信、同住或同机等多源交叉。", "计分方式": "按交叉来源累计，最高14分", "分值": 4, "上限": 14},

    # 资金代持/白手套：突出快进快出、双向对倒、净额小流水大、多账户、现金、整额和同日批量。
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-01", "规则说明": "高频大额往来：同一对象交易次数较多且累计绝对金额较高。", "计分方式": "按笔数和金额累计，最高16分", "分值": 2, "上限": 16},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-02", "规则说明": "双向资金对倒：同一对象既转入又转出，双向金额比例较高。", "计分方式": "按双向金额和比例累计，最高22分", "分值": 3, "上限": 22},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-03", "规则说明": "快进快出：7日内出现反向近似金额交易，疑似通道、过桥或代转。", "计分方式": "每组6分，最高30分", "分值": 6, "上限": 30},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-04", "规则说明": "现金关联：同一对象资金往来中存在现金、ATM、卡存、卡取等记录。", "计分方式": "每笔3分并叠加金额，最高18分", "分值": 3, "上限": 18},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-05", "规则说明": "代持/周转文本：摘要或备注含代持、代收、代付、周转、走账、过桥、备用金、借还款等。", "计分方式": "每笔5分，最高25分", "分值": 5, "上限": 25},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-06", "规则说明": "净额小但流水大：累计流水较大，但净额占流水比例较低，疑似资金中转。", "计分方式": "满足基础条件10分，金额较大继续叠加，最高20分", "分值": 10, "上限": 20},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-07", "规则说明": "多账户特征：同一对象关联多个对方账号或卡号。", "计分方式": "超过1个后每个4分，最高16分", "分值": 4, "上限": 16},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-08", "规则说明": "拆分/分散交易：同一对象多笔中等金额交易，累计金额较高。", "计分方式": "每笔2分，最高16分", "分值": 2, "上限": 16},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-09", "规则说明": "整额大额交易：万元级、千元级整额交易多次出现，疑似人为拆转或归集。", "计分方式": "每笔2分，最高12分", "分值": 2, "上限": 12},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-10", "规则说明": "同日批量交易：同一对象在同一天发生多笔资金往来或单日累计较高。", "计分方式": "按同日批量天数和单日笔数累计，最高14分", "分值": 2, "上限": 14},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-11", "规则说明": "高占比单一通道：单一对象占本人资金往来比例较高且有多笔往来。", "计分方式": "按占比和金额累计，最高12分", "分值": 2, "上限": 12},
    {"线索类型": "疑似资金代持/白手套线索", "规则编号": "WG-12", "规则说明": "现金叠加快进快出/对倒：现金记录与快进快出、双向对倒等同时出现。", "计分方式": "按组合特征累计，最高12分", "分值": 4, "上限": 12},

    # 利益输送：突出职务/机构对象、转入方向、敏感时点、利益文本、现金和通信叠加。
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-01", "规则说明": "职务对象转入：有职务信息的对象向查询对象转入资金。", "计分方式": "按转入金额和笔数累计，最高24分", "分值": 3, "上限": 24},
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-02", "规则说明": "机构对象转入：公司、工程、贸易、服务、投资等机构类对象向查询对象转入资金。", "计分方式": "按转入金额和笔数累计，最高24分", "分值": 3, "上限": 24},
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-03", "规则说明": "利益输送文本：摘要或备注含好处费、感谢费、辛苦费、回扣、协调费、服务费、居间费、项目款、工程款等。", "计分方式": "每笔5分，最高25分", "分值": 5, "上限": 25},
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-04", "规则说明": "敏感时点转入：非工作时间、节假日或深夜由职务/机构对象向查询对象转入资金。", "计分方式": "每笔4分，最高18分", "分值": 4, "上限": 18},
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-05", "规则说明": "拆分转入：同一对象多笔中等金额转入，累计金额较高。", "计分方式": "每笔3分，最高18分", "分值": 3, "上限": 18},
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-06", "规则说明": "长期稳定转入：同一对象跨月、多次向查询对象转入资金。", "计分方式": "按月份和笔数累计，最高16分", "分值": 2, "上限": 16},
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-07", "规则说明": "资金叠加通信：转入资金对象同时存在较高频通信或敏感时段通信。", "计分方式": "按通信强度累计，最高16分", "分值": 2, "上限": 16},
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-08", "规则说明": "职务/机构对象现金转入：职务或机构对象转入资金中出现现金、ATM、卡存等记录。", "计分方式": "每笔4分，最高16分", "分值": 4, "上限": 16},
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-09", "规则说明": "整额转入：职务或机构对象多次整额转入，金额形态异常规整。", "计分方式": "每笔2分，最高12分", "分值": 2, "上限": 12},
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-10", "规则说明": "高占比转入：职务或机构对象转入金额占该对象往来的主要部分。", "计分方式": "按转入占比和金额累计，最高12分", "分值": 2, "上限": 12},
    {"线索类型": "疑似利益输送线索", "规则编号": "INT-11", "规则说明": "项目/工程/服务类机构转入：机构名称或摘要呈项目、工程、服务、咨询、劳务等业务关联特征。", "计分方式": "按机构转入和文本叠加，最高14分", "分值": 3, "上限": 14},
]

PROFILE_NUMERIC_COLS = [
    "交易次数", "资金往来总额", "对象资金占比%", "转入笔数", "转入金额", "转入占比%", "转出笔数", "转出金额", "转出占比%", "净流入", "最大单笔", "最大单笔转入",
    "万元以上交易次数", "特殊金额次数", "情感关键词次数", "生活消费关键词次数", "亲密场景交易次数", "代持关键词次数", "利益输送关键词次数", "项目业务关键词次数",
    "非工作时间交易次数", "深夜交易次数", "特殊节日交易次数", "情感节日交易次数", "情感敏感日交易次数", "节假日交易次数", "敏感时点转入次数",
    "小额交易次数", "小额转入次数", "小额转出次数", "双向小额交易次数", "交易天数", "交易月份数", "转出月份数",
    "现金相关次数", "现金相关金额", "现金转入次数", "现金转出次数", "快进快出匹配组数", "快进快出涉及金额", "对方账户数", "对方卡号数",
    "拆分转入笔数", "拆分转入金额", "转入月份数", "整额大额交易次数", "整额大额转入次数", "同日多笔交易天数", "单日最大交易笔数", "单日最大往来金额",
    "通信次数", "非工作时间通信次数", "深夜通信次数(23–5)", "通话≥3分钟次数", "节日通信次数", "旅馆同住次数", "同机次数",
]


def _clue_rule_df() -> pd.DataFrame:
    d = pd.DataFrame(CLUE_RULE_DEFINITIONS)
    d.insert(0, "说明", "仅为自动化疑似线索筛查规则，不直接认定事实或性质；分值用于排序和提示，需结合原始凭证、身份关系、业务背景和询问核实。")
    return d


def _risk_rank(level: str) -> int:
    return {"无": 0, "低": 1, "中": 2, "高": 3}.get(safe_str(level).strip(), 0)


def _risk_grade(score: Any) -> str:
    """小分值累计后的风险分层。

    本版将单项规则分值调低，改为多因素叠加触发：
    1-17 分只作为低风险提示；18-34 分为中风险；35 分及以上为高风险。
    """
    try:
        s = float(score)
    except Exception:
        s = 0.0
    if s >= 35:
        return "高"
    if s >= 18:
        return "中"
    if s > 0:
        return "低"
    return "无"


def _is_hit_score(score: Any) -> bool:
    """可疑线索清单阈值。低于 12 分仍保留在画像表，不进入命中清单。"""
    try:
        return float(score) >= 12
    except Exception:
        return False


def _safe_int_num(v: Any) -> int:
    """将各种数值/空值安全转成 int。

    重点规避 Excel/通信统计中常见的 NaN、空字符串、inf 导致的
    “cannot convert float NaN to integer” 或非有限值转换错误。
    """
    try:
        vv = pd.to_numeric(v, errors="coerce")
        if isinstance(vv, pd.Series):
            vv = vv.iloc[0] if len(vv) else 0
        if pd.isna(vv):
            return 0
        try:
            if not np.isfinite(float(vv)):
                return 0
        except Exception:
            return 0
        return int(float(vv))
    except Exception:
        return 0


def _safe_float_num(v: Any) -> float:
    """将各种数值/空值安全转成 float，统一屏蔽 NaN/inf。"""
    try:
        vv = pd.to_numeric(v, errors="coerce")
        if isinstance(vv, pd.Series):
            vv = vv.iloc[0] if len(vv) else 0
        if pd.isna(vv):
            return 0.0
        f = float(vv)
        if not np.isfinite(f):
            return 0.0
        return f
    except Exception:
        return 0.0


def _safe_series_sum(series: Any) -> float:
    try:
        s = pd.to_numeric(series, errors="coerce").replace([np.inf, -np.inf], np.nan).fillna(0)
        return float(s.sum(skipna=True))
    except Exception:
        return 0.0


def _safe_series_max(series: Any) -> float:
    try:
        s = pd.to_numeric(series, errors="coerce").replace([np.inf, -np.inf], np.nan)
        if s.dropna().empty:
            return 0.0
        val = float(s.max(skipna=True))
        return val if np.isfinite(val) else 0.0
    except Exception:
        return 0.0


def _is_natural_person_name(name: Any) -> bool:
    s = safe_str(name).strip()
    return bool(CH_NAME_2_4_RE.fullmatch(s)) and not COMPANY_KEYWORDS_RE.search(s)


def _is_company_like(name: Any) -> bool:
    return bool(COMPANY_KEYWORDS_RE.search(safe_str(name).strip()))


def _is_special_emotion_amount(v: Any) -> bool:
    try:
        x = round(abs(float(v)), 2)
        if not np.isfinite(x):
            return False
    except Exception:
        return False
    if x in SPECIAL_EMOTION_AMOUNTS:
        return True
    # 兼容“1520、2521、11314、8999”等尾数表达，避免只识别固定金额。
    if abs(x - round(x)) < 0.005:
        s = str(int(round(x)))
        if len(s) >= 3 and any(s.endswith(tail) for tail in SPECIAL_EMOTION_AMOUNT_TAILS):
            return True
    return False


def _is_round_large_amount(v: Any) -> bool:
    """万元以上且金额呈整千/整万形态。该特征只作为累计因子，不单独定性。"""
    try:
        x = abs(float(v))
        if not np.isfinite(x):
            return False
    except Exception:
        return False
    if x < 10_000 or abs(x - round(x)) >= 0.005:
        return False
    xi = int(round(x))
    return xi % 10_000 == 0 or xi % 5_000 == 0 or xi % 1_000 == 0


def _is_romance_sensitive_date(ts: Any, festival: Any = "", text_blob: Any = "") -> bool:
    """识别情感关系常见敏感日期。

    包括 2/14、3/8、5/20、5/21、12/24、12/25，以及节日识别出的七夕；
    生日、纪念日等无法从日期本身判断，使用摘要/备注文本补充。
    """
    txt = safe_str(text_blob)
    if re.search(r"生日|纪念日|情人节|七夕|520|521", txt, flags=re.IGNORECASE):
        return True
    if safe_str(festival).strip() in {"七夕节"}:
        return True
    try:
        t = pd.to_datetime(ts, errors="coerce")
    except Exception:
        t = pd.NaT
    if pd.isna(t):
        return False
    return (int(t.month), int(t.day)) in {(2, 14), (3, 8), (5, 20), (5, 21), (12, 24), (12, 25)}


def _fmt_pct(v: Any) -> str:
    try:
        return f"{float(v):.1f}%"
    except Exception:
        return "0.0%"


def _unique_join(series: Any, sep: str = "、", limit: int = 8) -> str:
    vals: List[str] = []
    try:
        iterator = series.tolist()
    except Exception:
        iterator = list(series) if series is not None else []
    for v in iterator:
        s = safe_str(v).strip()
        if s and s.lower() != "nan" and s not in vals:
            vals.append(s)
        if limit and len(vals) >= limit:
            break
    return sep.join(vals)


def _mode_or_empty(series: Any) -> str:
    try:
        s = pd.Series(series).fillna("").map(safe_str).str.strip()
    except Exception:
        return ""
    s = s[(s != "") & (s.str.lower() != "nan")]
    if s.empty:
        return ""
    return s.value_counts().idxmax()


def _build_text_blob(d: pd.DataFrame) -> pd.Series:
    parts = []
    for c in CLUE_TEXT_COLS:
        if c in d.columns:
            parts.append(d[c].map(safe_str))
    if not parts:
        return pd.Series([""] * len(d), index=d.index, dtype=object)
    return pd.concat(parts, axis=1).agg(" ".join, axis=1).str.strip()


def _relation_object_from_row(row: pd.Series) -> str:
    name = safe_str(row.get("交易对方姓名", "")).strip()
    if name:
        return name
    for c in ["交易对方账户", "交易对方卡号", "交易对方证件号码"]:
        v = safe_str(row.get(c, "")).strip()
        if v:
            return f"账号/卡号:{v}"
    return "交易对手缺失"


def _prepare_clue_df(df: pd.DataFrame, person: str) -> pd.DataFrame:
    d = df.copy() if df is not None else pd.DataFrame()
    need = [
        "查询对象", "交易对方姓名", "对方职务", "借贷标志", "交易金额", "交易时间", "交易流水号", "节假日", "星期",
        "交易对方账户", "交易对方卡号", "交易对方证件号码", "交易对方账号开户行", "反馈单位", "查询账户", "查询卡号", "来源文件",
    ] + CLUE_TEXT_COLS
    for c in need:
        if c not in d.columns:
            d[c] = ""
    d["查询对象"] = d["查询对象"].map(safe_str).str.strip().replace("", person)
    d["交易对方姓名"] = d["交易对方姓名"].map(safe_str).str.strip()
    d["对方职务"] = d["对方职务"].map(safe_str).str.strip()
    d["交易时间"] = pd.to_datetime(d["交易时间"], errors="coerce")
    d["__amt__"] = pd.to_numeric(d["交易金额"], errors="coerce").replace([np.inf, -np.inf], np.nan).fillna(0.0)
    d["__abs_amt__"] = d["__amt__"].abs().fillna(0.0)

    is_in = d["借贷标志"].map(safe_str).str.strip().eq("进")
    is_out = d["借贷标志"].map(safe_str).str.strip().eq("出")
    fallback_in = (~is_in) & (~is_out) & (d["__amt__"] > 0)
    fallback_out = (~is_in) & (~is_out) & (d["__amt__"] < 0)
    d["__is_in__"] = is_in | fallback_in
    d["__is_out__"] = is_out | fallback_out

    directions = []
    for _, r in d.iterrows():
        obj = safe_str(r.get("交易对方姓名", "")).strip() or "交易对手"
        if bool(r.get("__is_in__", False)):
            directions.append(f"{obj} → {person}")
        elif bool(r.get("__is_out__", False)):
            directions.append(f"{person} → {obj}")
        else:
            directions.append("方向未明")
    d["__方向说明__"] = directions

    d["关系对象"] = d.apply(_relation_object_from_row, axis=1)
    d["__festival__"] = _festival_series(d["交易时间"])
    d["__offwork__"] = _flag_offwork(d["交易时间"]).fillna(False)
    d["__late_night__"] = _flag_late_night(d["交易时间"]).fillna(False)
    d["__natural_person__"] = d["交易对方姓名"].map(_is_natural_person_name)
    d["__company_like__"] = d["交易对方姓名"].map(_is_company_like)
    d["__text_blob__"] = _build_text_blob(d)
    d["__emotion_keyword__"] = d["__text_blob__"].str.contains(EMOTION_KEYWORDS_RE, na=False)
    d["__lifecare_keyword__"] = d["__text_blob__"].str.contains(LIFECARE_KEYWORDS_RE, na=False)
    d["__holding_keyword__"] = d["__text_blob__"].str.contains(HOLDING_KEYWORDS_RE, na=False)
    d["__benefit_keyword__"] = d["__text_blob__"].str.contains(BENEFIT_KEYWORDS_RE, na=False)
    d["__project_keyword__"] = d["__text_blob__"].str.contains(PROJECT_KEYWORDS_RE, na=False)
    d["__special_emotion_amount__"] = d["__abs_amt__"].map(_is_special_emotion_amount)
    d["__round_big_amt__"] = d["__abs_amt__"].map(_is_round_large_amount)
    d["__romance_day__"] = [
        _is_romance_sensitive_date(ts, fest, txt)
        for ts, fest, txt in zip(d["交易时间"], d["__festival__"], d["__text_blob__"])
    ]
    d["__cash_signal__"] = _cash_signal_mask(d).reindex(d.index, fill_value=False)
    d["__small_amt__"] = (d["__abs_amt__"] > 0) & (d["__abs_amt__"] <= 5_000)
    d["__mid_amt__"] = (d["__abs_amt__"] >= 5_000) & (d["__abs_amt__"] <= 50_000)
    d["__month__"] = d["交易时间"].dt.to_period("M").astype(str).replace("NaT", "")
    d["__date__"] = d["交易时间"].dt.date
    return d


def _quick_pass_stats(g: pd.DataFrame, max_days: int = 7, min_amount: float = 10_000, diff_rate_limit: float = 0.05) -> Tuple[int, float, str, set]:
    if g is None or g.empty:
        return 0, 0.0, "", set()
    gg = g.dropna(subset=["交易时间", "__abs_amt__"]).copy()
    if len(gg) < 2:
        return 0, 0.0, "", set()
    in_rows = gg[gg["__is_in__"] & (gg["__abs_amt__"] >= min_amount)]
    out_rows = gg[gg["__is_out__"] & (gg["__abs_amt__"] >= min_amount)]
    if in_rows.empty or out_rows.empty:
        return 0, 0.0, "", set()

    pairs: List[str] = []
    matched_idx: set = set()
    total_amt = 0.0
    max_delta = pd.Timedelta(days=max_days)
    for i, rin in in_rows.iterrows():
        amt = _safe_float_num(rin.get("__abs_amt__", 0))
        if amt <= 0:
            continue
        candidates = out_rows[out_rows["交易时间"].sub(rin["交易时间"]).abs() <= max_delta]
        if candidates.empty:
            continue
        diff_rate = candidates["__abs_amt__"].sub(amt).abs() / amt
        near = candidates[diff_rate <= diff_rate_limit]
        if near.empty:
            continue
        best = near.assign(__diff__=diff_rate.loc[near.index]).sort_values("__diff__", kind="mergesort").iloc[0]
        matched_idx.add(i)
        matched_idx.add(best.name)
        total_amt += amt
        if len(pairs) < 5:
            d1 = rin["交易时间"].strftime("%Y-%m-%d") if pd.notna(rin["交易时间"]) else "日期缺失"
            d2 = best["交易时间"].strftime("%Y-%m-%d") if pd.notna(best["交易时间"]) else "日期缺失"
            pairs.append(f"{d1}转入{_fmt_money_human(amt)}，{d2}反向{_fmt_money_human(_safe_float_num(best.get('__abs_amt__', 0)))}")
    return len(pairs), total_amt, "；".join(pairs), matched_idx


def build_bank_counterparty_profile(txn_df: pd.DataFrame, person: str) -> pd.DataFrame:
    d = _prepare_clue_df(txn_df, person)
    if d.empty:
        return pd.DataFrame()

    total_person_abs = _safe_series_sum(d["__abs_amt__"])
    rows: List[Dict[str, Any]] = []
    for obj, g in d.groupby("关系对象", dropna=False):
        obj = safe_str(obj).strip() or "交易对手缺失"
        if obj == "交易对手缺失" and g["__abs_amt__"].fillna(0).sum() == 0:
            continue
        g = g.copy()
        in_g = g[g["__is_in__"]]
        out_g = g[g["__is_out__"]]
        quick_cnt, quick_amt, quick_sample, matched_idx = _quick_pass_stats(g)

        total_abs = _safe_series_sum(g["__abs_amt__"])
        in_sum = _safe_series_sum(in_g["__abs_amt__"])
        out_sum = _safe_series_sum(out_g["__abs_amt__"])
        net = in_sum - out_sum
        max_amt = _safe_series_max(g["__abs_amt__"])
        max_in_amt = _safe_series_max(in_g["__abs_amt__"]) if not in_g.empty else 0.0
        name = _mode_or_empty(g["交易对方姓名"])
        title = _mode_or_empty(g["对方职务"])
        first_dt = g["交易时间"].min() if g["交易时间"].notna().any() else pd.NaT
        last_dt = g["交易时间"].max() if g["交易时间"].notna().any() else pd.NaT
        acct_ser = g["交易对方账户"].map(safe_str).str.strip()
        card_ser = g["交易对方卡号"].map(safe_str).str.strip()
        acct_count = acct_ser.mask(acct_ser == "").nunique(dropna=True)
        card_count = card_ser.mask(card_ser == "").nunique(dropna=True)
        split_in = in_g[(in_g["__abs_amt__"] >= 5_000) & (in_g["__abs_amt__"] <= 50_000)]
        cash_g = g[g["__cash_signal__"]]
        cash_in = in_g[in_g["__cash_signal__"]]
        cash_out = out_g[out_g["__cash_signal__"]]
        small_g = g[g["__small_amt__"]]
        small_in = in_g[in_g["__small_amt__"]]
        small_out = out_g[out_g["__small_amt__"]]
        sensitive_in = in_g[in_g["__offwork__"] | in_g["__late_night__"] | in_g["__festival__"].map(safe_str).str.strip().ne("") | in_g["节假日"].map(safe_str).str.strip().eq("节假日")]
        valid_dates = g["__date__"].dropna()
        tx_days = int(pd.Series(valid_dates).nunique()) if len(valid_dates) else 0
        tx_months = int(g["__month__"].replace("", np.nan).nunique(dropna=True))
        out_months = int(out_g["__month__"].replace("", np.nan).nunique(dropna=True)) if not out_g.empty else 0
        in_months = int(in_g["__month__"].replace("", np.nan).nunique(dropna=True)) if not in_g.empty else 0
        if valid_dates.empty:
            same_day_days = 0
            same_day_max_cnt = 0
            same_day_max_amt = 0.0
        else:
            day_grp = g.dropna(subset=["__date__"]).groupby("__date__", dropna=False)
            day_cnt = day_grp.size()
            day_amt = day_grp["__abs_amt__"].sum()
            same_day_days = int((day_cnt >= 3).sum())
            same_day_max_cnt = _safe_int_num(day_cnt.max()) if not day_cnt.empty else 0
            same_day_max_amt = _safe_series_max(day_amt) if not day_amt.empty else 0.0
        row = {
            "查询对象": person,
            "关系对象": obj,
            "交易对方姓名": name,
            "对象职务": title,
            "是否自然人对象": "是" if bool(g["__natural_person__"].any()) else "",
            "是否机构类对象": "是" if bool(g["__company_like__"].any()) else "",
            "对象账号": _unique_join(g["交易对方账户"], limit=10),
            "对象卡号": _unique_join(g["交易对方卡号"], limit=10),
            "对象开户行": _unique_join(g["交易对方账号开户行"], limit=5),
            "交易次数": int(len(g)),
            "资金往来总额": total_abs,
            "对象资金占比%": (total_abs / total_person_abs * 100) if total_person_abs else 0,
            "转入笔数": int(len(in_g)),
            "转入金额": in_sum,
            "转入占比%": (in_sum / total_abs * 100) if total_abs else 0,
            "转出笔数": int(len(out_g)),
            "转出金额": out_sum,
            "转出占比%": (out_sum / total_abs * 100) if total_abs else 0,
            "净流入": net,
            "最大单笔": max_amt,
            "最大单笔转入": max_in_amt,
            "万元以上交易次数": int((g["__abs_amt__"] >= 10_000).sum()),
            "特殊金额次数": int(g["__special_emotion_amount__"].sum()),
            "情感关键词次数": int(g["__emotion_keyword__"].sum()),
            "生活消费关键词次数": int(g["__lifecare_keyword__"].sum()),
            "亲密场景交易次数": int((g["__emotion_keyword__"] | g["__lifecare_keyword__"] | g["__romance_day__"]).sum()),
            "代持关键词次数": int(g["__holding_keyword__"].sum()),
            "利益输送关键词次数": int(g["__benefit_keyword__"].sum()),
            "项目业务关键词次数": int(g["__project_keyword__"].sum()),
            "非工作时间交易次数": int(g["__offwork__"].sum()),
            "深夜交易次数": int(g["__late_night__"].sum()),
            "特殊节日交易次数": int(g["__festival__"].map(safe_str).str.strip().ne("").sum()),
            "情感节日交易次数": int(g["__festival__"].isin(["七夕节", "5月20日"]).sum()),
            "情感敏感日交易次数": int(pd.Series(g["__romance_day__"], index=g.index).sum()),
            "节假日交易次数": int(g["节假日"].map(safe_str).str.strip().eq("节假日").sum()),
            "敏感时点转入次数": int(len(sensitive_in)),
            "小额交易次数": int(len(small_g)),
            "小额转入次数": int(len(small_in)),
            "小额转出次数": int(len(small_out)),
            "双向小额交易次数": int(len(small_in) + len(small_out)) if (len(small_in) > 0 and len(small_out) > 0) else 0,
            "交易天数": tx_days,
            "交易月份数": tx_months,
            "转出月份数": out_months,
            "现金相关次数": int(len(cash_g)),
            "现金相关金额": _safe_series_sum(cash_g["__abs_amt__"]),
            "现金转入次数": int(len(cash_in)),
            "现金转出次数": int(len(cash_out)),
            "快进快出匹配组数": int(quick_cnt),
            "快进快出涉及金额": float(quick_amt),
            "快进快出样例": quick_sample,
            "快进快出流水索引": "、".join(map(str, sorted(matched_idx))) if matched_idx else "",
            "对方账户数": int(acct_count),
            "对方卡号数": int(card_count),
            "拆分转入笔数": int(len(split_in)),
            "拆分转入金额": _safe_series_sum(split_in["__abs_amt__"]),
            "转入月份数": in_months,
            "整额大额交易次数": int(g["__round_big_amt__"].sum()),
            "整额大额转入次数": int(in_g["__round_big_amt__"].sum()) if not in_g.empty else 0,
            "同日多笔交易天数": same_day_days,
            "单日最大交易笔数": same_day_max_cnt,
            "单日最大往来金额": same_day_max_amt,
            "首次交易时间": first_dt.strftime("%Y-%m-%d %H:%M:%S") if pd.notna(first_dt) else "",
            "末次交易时间": last_dt.strftime("%Y-%m-%d %H:%M:%S") if pd.notna(last_dt) else "",
        }
        rows.append(row)
    return pd.DataFrame(rows).replace({np.nan: ""})


def build_comm_counterparty_profile(person: str) -> pd.DataFrame:
    st = STATE.COMM_STATS_BY_PERSON.get(person)
    if st is None or st.empty:
        return pd.DataFrame()
    d = st.copy().replace({np.nan: ""})
    for c in ["姓名", "职务", "对方号码", "通信次数", "非工作时间通信次数", "深夜通信次数(23–5)", "通话≥3分钟次数"]:
        if c not in d.columns:
            d[c] = ""
    fest_cols = [f"{x}通信次数" for x in FESTIVAL_NAMES]
    for c in fest_cols:
        if c not in d.columns:
            d[c] = 0
    d["关系对象"] = np.where(d["姓名"].map(safe_str).str.strip().ne(""), d["姓名"].map(safe_str).str.strip(), d["对方号码"].map(safe_str).str.strip())
    d = d[d["关系对象"].map(safe_str).str.strip().ne("")].copy()
    if d.empty:
        return pd.DataFrame()

    rows: List[Dict[str, Any]] = []
    for obj, g in d.groupby("关系对象", dropna=False):
        fest_total = 0
        for c in fest_cols:
            fest_total += _safe_int_num(_safe_series_sum(g[c]))
        rows.append({
            "查询对象": person,
            "关系对象": safe_str(obj).strip(),
            "交易对方姓名": _mode_or_empty(g["姓名"]),
            "对象职务": _mode_or_empty(g["职务"]),
            "对象手机号": _unique_join(g["对方号码"], limit=10),
            "通信次数": _safe_int_num(_safe_series_sum(g["通信次数"])),
            "非工作时间通信次数": _safe_int_num(_safe_series_sum(g["非工作时间通信次数"])),
            "深夜通信次数(23–5)": _safe_int_num(_safe_series_sum(g["深夜通信次数(23–5)"])),
            "通话≥3分钟次数": _safe_int_num(_safe_series_sum(g["通话≥3分钟次数"])),
            "节日通信次数": _safe_int_num(fest_total),
        })
    return pd.DataFrame(rows).replace({np.nan: ""})


def build_co_travel_profile(person: str) -> pd.DataFrame:
    """读取已标注的旅馆同住/同机结果，补充到关系对象画像。"""
    try:
        root = out_gongan_dir(person)
    except Exception:
        return pd.DataFrame()
    if not root.exists():
        return pd.DataFrame()

    rows: List[Dict[str, Any]] = []
    for p in root.glob("*.xlsx"):
        stem = p.stem
        if "旅馆同住" not in stem and "同机" not in stem:
            continue
        kind = "旅馆同住" if "旅馆同住" in stem else "同机"
        base_col = HOTEL_COL_NAME if kind == "旅馆同住" else FLIGHT_COL_NAME
        try:
            df = pd.read_excel(p)
        except Exception:
            continue
        if df is None or df.empty:
            continue
        df = clean_columns(df)
        if base_col not in df.columns:
            continue
        for _, r in df.iterrows():
            obj = safe_str(r.get(base_col, "")).strip()
            if not obj or obj == person:
                continue
            rows.append({
                "查询对象": person,
                "关系对象": obj,
                "交易对方姓名": obj,
                "对象职务": safe_str(r.get("职务", "")).strip(),
                "旅馆同住次数": 1 if kind == "旅馆同住" else 0,
                "同机次数": 1 if kind == "同机" else 0,
            })
    if not rows:
        return pd.DataFrame()
    d = pd.DataFrame(rows)
    out_rows: List[Dict[str, Any]] = []
    for obj, g in d.groupby("关系对象", dropna=False):
        out_rows.append({
            "查询对象": person,
            "关系对象": safe_str(obj).strip(),
            "交易对方姓名": _mode_or_empty(g["交易对方姓名"]),
            "对象职务": _mode_or_empty(g["对象职务"]),
            "旅馆同住次数": _safe_int_num(_safe_series_sum(g["旅馆同住次数"])),
            "同机次数": _safe_int_num(_safe_series_sum(g["同机次数"])),
        })
    return pd.DataFrame(out_rows).replace({np.nan: ""})


def merge_relation_profiles(bank_profile: pd.DataFrame, comm_profile: pd.DataFrame, co_travel_profile: pd.DataFrame, person: str) -> pd.DataFrame:
    rows: Dict[str, Dict[str, Any]] = {}

    def ensure_row(obj: Any) -> Dict[str, Any]:
        key = safe_str(obj).strip() or "交易对手缺失"
        if key not in rows:
            rows[key] = {"查询对象": person, "关系对象": key}
            for c in PROFILE_NUMERIC_COLS:
                rows[key][c] = 0
        return rows[key]

    def merge_one(df: pd.DataFrame) -> None:
        if df is None or df.empty or "关系对象" not in df.columns:
            return
        for _, r in df.iterrows():
            obj = safe_str(r.get("关系对象", "")).strip()
            if not obj:
                continue
            base = ensure_row(obj)
            for c, v in r.items():
                if c in ["查询对象", "关系对象"]:
                    continue
                if c in PROFILE_NUMERIC_COLS:
                    base[c] = _safe_float_num(base.get(c, 0)) + _safe_float_num(v)
                else:
                    old = safe_str(base.get(c, "")).strip()
                    new_v = safe_str(v).strip()
                    if not old and new_v:
                        base[c] = new_v
                    elif old and new_v and c in {"对象手机号", "对象账号", "对象卡号", "对象开户行"}:
                        vals = []
                        for part in (old + "、" + new_v).split("、"):
                            part = part.strip()
                            if part and part not in vals:
                                vals.append(part)
                        base[c] = "、".join(vals[:10])

    merge_one(bank_profile)
    merge_one(comm_profile)
    merge_one(co_travel_profile)

    if not rows:
        return pd.DataFrame()
    profile = pd.DataFrame(rows.values()).replace({np.nan: ""})
    for c in PROFILE_NUMERIC_COLS:
        if c not in profile.columns:
            profile[c] = 0
        profile[c] = pd.to_numeric(profile[c], errors="coerce").fillna(0)

    if "对象职务" not in profile.columns:
        profile["对象职务"] = ""
    if "交易对方姓名" not in profile.columns:
        profile["交易对方姓名"] = ""
    natural_existing = (
        profile["是否自然人对象"].astype(str).str.strip().eq("是")
        if "是否自然人对象" in profile.columns
        else pd.Series(False, index=profile.index)
    )
    company_existing = (
        profile["是否机构类对象"].astype(str).str.strip().eq("是")
        if "是否机构类对象" in profile.columns
        else pd.Series(False, index=profile.index)
    )
    profile["是否自然人对象"] = np.where(natural_existing | profile["关系对象"].map(_is_natural_person_name), "是", "")
    profile["是否机构类对象"] = np.where(company_existing | profile["关系对象"].map(_is_company_like), "是", "")
    profile["平均单笔金额"] = np.where(profile["交易次数"] > 0, profile["资金往来总额"] / profile["交易次数"], 0)
    profile["双向金额比例"] = np.where(
        np.maximum(profile["转入金额"], profile["转出金额"]) > 0,
        np.minimum(profile["转入金额"], profile["转出金额"]) / np.maximum(profile["转入金额"], profile["转出金额"]),
        0,
    )
    profile["净额占流水比例"] = np.where(profile["资金往来总额"] > 0, profile["净流入"].abs() / profile["资金往来总额"], 0)
    sort_cols = ["资金往来总额", "通信次数", "旅馆同住次数", "同机次数"]
    profile.sort_values(sort_cols, ascending=[False, False, False, False], inplace=True, kind="mergesort")
    return profile.reset_index(drop=True)


def _profile_metrics_sentence(row: pd.Series) -> str:
    """把对象画像压缩成一行可读证据。"""
    parts = []
    tx_cnt = _safe_int_num(row.get("交易次数", 0))
    if tx_cnt:
        parts.append(f"资金{tx_cnt}笔/{_fmt_money_human(row.get('资金往来总额', 0))}")
        parts.append(f"入{_safe_int_num(row.get('转入笔数', 0))}笔{_fmt_money_human(row.get('转入金额', 0))}")
        parts.append(f"出{_safe_int_num(row.get('转出笔数', 0))}笔{_fmt_money_human(row.get('转出金额', 0))}")
        parts.append(f"最大{_fmt_money_human(row.get('最大单笔', 0))}")
    key_metrics = [
        ("特殊金额", "特殊金额次数", "笔"),
        ("情感备注", "情感关键词次数", "笔"),
        ("生活类备注", "生活消费关键词次数", "笔"),
        ("情感敏感日", "情感敏感日交易次数", "笔"),
        ("小额", "小额交易次数", "笔"),
        ("整额大额", "整额大额交易次数", "笔"),
        ("同日批量", "同日多笔交易天数", "天"),
        ("现金", "现金相关次数", "笔"),
        ("快进快出", "快进快出匹配组数", "组"),
        ("代持文本", "代持关键词次数", "笔"),
        ("利益文本", "利益输送关键词次数", "笔"),
        ("项目/工程文本", "项目业务关键词次数", "笔"),
    ]
    for label, col, unit in key_metrics:
        v = _safe_int_num(row.get(col, 0))
        if v:
            parts.append(f"{label}{v}{unit}")
    months = _safe_int_num(row.get("交易月份数", 0))
    if months:
        parts.append(f"跨{months}个月")
    share = _safe_float_num(row.get("对象资金占比%", 0))
    if share >= 10:
        parts.append(f"占比{share:.1f}%")
    comm_cnt = _safe_int_num(row.get("通信次数", 0))
    if comm_cnt:
        parts.append(f"通信{comm_cnt}次")
    for label, col in [("深夜通信", "深夜通信次数(23–5)"), ("节日通信", "节日通信次数"), ("长通话", "通话≥3分钟次数")]:
        v = _safe_int_num(row.get(col, 0))
        if v:
            parts.append(f"{label}{v}次")
    hotel = _safe_int_num(row.get("旅馆同住次数", 0))
    flight = _safe_int_num(row.get("同机次数", 0))
    if hotel:
        parts.append(f"同住{hotel}次")
    if flight:
        parts.append(f"同机{flight}次")
    return "；".join(parts[:14])


def _score_relation_row(row: pd.Series) -> Dict[str, List[Dict[str, Any]]]:
    """对单个关系对象进行三类线索累计评分。

    设计原则：单项规则分值降低，单一异常只提示；多个规则、多源证据叠加后才进入可疑线索清单。
    """
    rules: Dict[str, List[Dict[str, Any]]] = {t: [] for t in CLUE_TYPES}

    def add(clue_type: str, rule_id: str, name: str, score: Any, evidence: str) -> None:
        score_i = max(0, int(round(_safe_float_num(score))))
        if score_i <= 0:
            return
        rules[clue_type].append({"规则编号": rule_id, "规则名称": name, "分值": score_i, "证据": evidence})

    is_person = safe_str(row.get("是否自然人对象", "")).strip() == "是"
    is_company = safe_str(row.get("是否机构类对象", "")).strip() == "是"
    has_title = bool(safe_str(row.get("对象职务", "")).strip())
    tx_cnt = _safe_int_num(row.get("交易次数", 0))
    total_abs = _safe_float_num(row.get("资金往来总额", 0))
    avg_amt = _safe_float_num(row.get("平均单笔金额", 0))
    in_cnt = _safe_int_num(row.get("转入笔数", 0))
    out_cnt = _safe_int_num(row.get("转出笔数", 0))
    in_sum = _safe_float_num(row.get("转入金额", 0))
    out_sum = _safe_float_num(row.get("转出金额", 0))
    in_ratio_pct = _safe_float_num(row.get("转入占比%", 0))
    out_ratio_pct = _safe_float_num(row.get("转出占比%", 0))
    max_in_amt = _safe_float_num(row.get("最大单笔转入", 0))
    ratio = _safe_float_num(row.get("双向金额比例", 0))
    net_ratio = _safe_float_num(row.get("净额占流水比例", 0))
    comm_cnt = _safe_int_num(row.get("通信次数", 0))
    comm_night = _safe_int_num(row.get("深夜通信次数(23–5)", 0))
    comm_fest = _safe_int_num(row.get("节日通信次数", 0))
    long_call = _safe_int_num(row.get("通话≥3分钟次数", 0))
    small_cnt = _safe_int_num(row.get("小额交易次数", 0))
    small_in = _safe_int_num(row.get("小额转入次数", 0))
    small_out = _safe_int_num(row.get("小额转出次数", 0))
    months = _safe_int_num(row.get("交易月份数", 0))
    in_months = _safe_int_num(row.get("转入月份数", 0))
    out_months = _safe_int_num(row.get("转出月份数", 0))
    obj_share = _safe_float_num(row.get("对象资金占比%", 0))
    special_amt_cnt = _safe_int_num(row.get("特殊金额次数", 0))
    emo_kw_cnt = _safe_int_num(row.get("情感关键词次数", 0))
    life_kw_cnt = _safe_int_num(row.get("生活消费关键词次数", 0))
    intimate_scene_cnt = _safe_int_num(row.get("亲密场景交易次数", 0))
    romance_day_cnt = _safe_int_num(row.get("情感敏感日交易次数", 0))
    holding_kw = _safe_int_num(row.get("代持关键词次数", 0))
    benefit_kw = _safe_int_num(row.get("利益输送关键词次数", 0))
    project_kw = _safe_int_num(row.get("项目业务关键词次数", 0))
    hotel_cnt = _safe_int_num(row.get("旅馆同住次数", 0))
    flight_cnt = _safe_int_num(row.get("同机次数", 0))
    offwork = _safe_int_num(row.get("非工作时间交易次数", 0))
    late = _safe_int_num(row.get("深夜交易次数", 0))
    holiday = _safe_int_num(row.get("特殊节日交易次数", 0)) + _safe_int_num(row.get("节假日交易次数", 0))
    quick_cnt = _safe_int_num(row.get("快进快出匹配组数", 0))
    cash_cnt = _safe_int_num(row.get("现金相关次数", 0))
    cash_amt = _safe_float_num(row.get("现金相关金额", 0))
    cash_in_cnt = _safe_int_num(row.get("现金转入次数", 0))
    cash_out_cnt = _safe_int_num(row.get("现金转出次数", 0))
    acct_total = _safe_int_num(row.get("对方账户数", 0)) + _safe_int_num(row.get("对方卡号数", 0))
    split_cnt = _safe_int_num(row.get("拆分转入笔数", 0))
    split_amt = _safe_float_num(row.get("拆分转入金额", 0))
    round_cnt = _safe_int_num(row.get("整额大额交易次数", 0))
    round_in_cnt = _safe_int_num(row.get("整额大额转入次数", 0))
    same_day_days = _safe_int_num(row.get("同日多笔交易天数", 0))
    same_day_max_cnt = _safe_int_num(row.get("单日最大交易笔数", 0))
    same_day_max_amt = _safe_float_num(row.get("单日最大往来金额", 0))

    # 一、疑似情感关系线索：重视自然人、情感表达、小额高频、非工作时间、通信和轨迹交叉。
    if special_amt_cnt > 0:
        add("疑似情感关系线索", "EMO-01", "特殊表达金额", min(12, special_amt_cnt * 2), f"特殊金额{special_amt_cnt}笔")
    if emo_kw_cnt > 0:
        add("疑似情感关系线索", "EMO-02", "亲密/情感文本", min(20, emo_kw_cnt * 4), f"亲密或情感备注{emo_kw_cnt}笔")
    if (is_person or emo_kw_cnt > 0) and romance_day_cnt > 0:
        add("疑似情感关系线索", "EMO-03", "情感敏感日期", min(15, romance_day_cnt * 3), f"情感敏感日期/文本交易{romance_day_cnt}笔")
    if is_person and tx_cnt >= 2 and (offwork + late + holiday) >= 2:
        st_score = min(14, offwork * 1 + late * 2 + holiday * 1)
        add("疑似情感关系线索", "EMO-04", "敏感时段资金往来", st_score, f"非工作{offwork}笔、深夜{late}笔、节日/节假日{holiday}笔")
    if is_person and (comm_cnt >= 3 or comm_night > 0 or comm_fest > 0 or long_call >= 2):
        comm_score = min(18, max(1, comm_cnt // 3) + comm_night * 2 + comm_fest * 3 + long_call // 2)
        add("疑似情感关系线索", "EMO-05", "资金叠加通信", comm_score, f"通信{comm_cnt}次、深夜{comm_night}次、节日{comm_fest}次、长通话{long_call}次")
    if hotel_cnt > 0:
        add("疑似情感关系线索", "EMO-06", "旅馆同住", min(36, hotel_cnt * 18), f"旅馆同住{hotel_cnt}次")
    if flight_cnt > 0:
        add("疑似情感关系线索", "EMO-07", "同机出行", min(16, flight_cnt * 8), f"同机{flight_cnt}次")
    if is_person and small_cnt >= 4:
        add("疑似情感关系线索", "EMO-08", "小额高频", min(12, small_cnt - 3), f"小额交易{small_cnt}笔，平均{_fmt_money_human(avg_amt)}")
    if is_person and small_in > 0 and small_out > 0:
        add("疑似情感关系线索", "EMO-09", "双向小额互转", min(12, min(small_in, small_out) * 2 + (small_in + small_out) // 5), f"小额转入{small_in}笔、小额转出{small_out}笔")
    if is_person and months >= 3 and tx_cnt >= 3:
        add("疑似情感关系线索", "EMO-10", "长期持续往来", min(12, (months - 2) * 3 + min(4, tx_cnt // 6)), f"跨{months}个月、交易{tx_cnt}笔")
    if is_person and life_kw_cnt > 0:
        add("疑似情感关系线索", "EMO-11", "生活照顾类文本", min(16, life_kw_cnt * 4), f"生活消费/照顾类备注{life_kw_cnt}笔")
    if is_person and out_sum >= 5_000 and out_cnt >= 2 and out_ratio_pct >= 70:
        score = min(14, 3 + out_cnt + int(out_sum // 20_000) + max(0, out_months - 1))
        add("疑似情感关系线索", "EMO-12", "单向供养/支持", score, f"转出{out_cnt}笔{_fmt_money_human(out_sum)}，转出占比{_fmt_pct(out_ratio_pct)}")
    if is_person and obj_share >= 15 and total_abs >= 10_000:
        add("疑似情感关系线索", "EMO-13", "自然人资金占比较高", min(10, int(obj_share // 10) * 2 + (2 if total_abs >= 50_000 else 0)), f"该自然人占本人资金往来{obj_share:.1f}%")
    if is_person and (special_amt_cnt or romance_day_cnt or intimate_scene_cnt):
        combo = 0
        combo += 1 if special_amt_cnt and (emo_kw_cnt or life_kw_cnt) else 0
        combo += 1 if special_amt_cnt and (offwork or late or holiday) else 0
        combo += 1 if romance_day_cnt and (offwork or late or holiday) else 0
        combo += 1 if intimate_scene_cnt and small_cnt >= 3 else 0
        combo += 1 if emo_kw_cnt and comm_cnt >= 3 else 0
        if combo:
            add("疑似情感关系线索", "EMO-14", "情感组合特征", min(12, combo * 2 + min(4, special_amt_cnt + romance_day_cnt)), f"特殊金额/亲密文本/敏感日期/敏感时段组合{combo}类")
    cross_sources = int(tx_cnt > 0) + int(comm_cnt > 0) + int(hotel_cnt > 0) + int(flight_cnt > 0)
    if is_person and cross_sources >= 2:
        score = min(14, (cross_sources - 1) * 4 + (4 if hotel_cnt else 0) + (2 if flight_cnt else 0))
        add("疑似情感关系线索", "EMO-15", "轨迹/通信/资金交叉", score, f"资金/通信/同住/同机交叉来源{cross_sources}类")

    # 二、疑似资金代持/白手套线索：重视双向、快进快出、净额小流水大、多账户、现金和同日批量。
    if tx_cnt >= 3 and total_abs >= 30_000:
        add("疑似资金代持/白手套线索", "WG-01", "高频大额往来", min(16, (tx_cnt // 2) * 2 + int(total_abs // 100_000) * 3), f"交易{tx_cnt}笔，累计{_fmt_money_human(total_abs)}")
    if min(in_sum, out_sum) >= 10_000 and ratio >= 0.35:
        add("疑似资金代持/白手套线索", "WG-02", "双向资金对倒", min(22, 5 + int(ratio * 10) + int(min(in_sum, out_sum) // 50_000) * 2), f"转入{_fmt_money_human(in_sum)}、转出{_fmt_money_human(out_sum)}，双向比例{ratio:.1%}")
    if quick_cnt > 0:
        sample = safe_str(row.get("快进快出样例", "")).strip()
        add("疑似资金代持/白手套线索", "WG-03", "7日内反向近似金额快进快出", min(30, quick_cnt * 6), sample or f"匹配{quick_cnt}组")
    if cash_cnt > 0 or cash_amt >= 20_000:
        add("疑似资金代持/白手套线索", "WG-04", "现金/ATM/卡存卡取关联", min(18, cash_cnt * 3 + int(cash_amt // 50_000) * 2), f"现金相关{cash_cnt}笔，金额{_fmt_money_human(cash_amt)}")
    if holding_kw > 0:
        add("疑似资金代持/白手套线索", "WG-05", "代持/周转/走账文本", min(25, holding_kw * 5), f"代持/周转/走账类关键词{holding_kw}笔")
    if total_abs >= 50_000 and net_ratio <= 0.30 and in_sum > 0 and out_sum > 0:
        add("疑似资金代持/白手套线索", "WG-06", "净额小但流水大", min(20, 10 + int(total_abs // 200_000) * 2 + (2 if net_ratio <= 0.15 else 0)), f"累计{_fmt_money_human(total_abs)}，净额占流水{net_ratio:.1%}")
    if acct_total >= 2:
        add("疑似资金代持/白手套线索", "WG-07", "多账户/多卡特征", min(16, (acct_total - 1) * 4), f"关联账号/卡号合计{acct_total}个")
    if split_cnt >= 3 and split_amt >= 30_000:
        add("疑似资金代持/白手套线索", "WG-08", "拆分/分散交易", min(16, split_cnt * 2 + int(split_amt // 100_000)), f"5千至5万元区间转入{split_cnt}笔，累计{_fmt_money_human(split_amt)}")
    if round_cnt >= 3 or (round_cnt >= 2 and total_abs >= 50_000):
        add("疑似资金代持/白手套线索", "WG-09", "整额大额交易", min(12, round_cnt * 2 + int(total_abs // 200_000)), f"万元以上整额/整千交易{round_cnt}笔")
    if same_day_days >= 1 and same_day_max_cnt >= 3:
        add("疑似资金代持/白手套线索", "WG-10", "同日批量交易", min(14, same_day_days * 3 + max(0, same_day_max_cnt - 2) * 2 + int(same_day_max_amt // 50_000)), f"同日多笔{same_day_days}天，单日最高{same_day_max_cnt}笔/{_fmt_money_human(same_day_max_amt)}")
    if obj_share >= 25 and tx_cnt >= 3 and total_abs >= 50_000:
        add("疑似资金代持/白手套线索", "WG-11", "高占比单一通道", min(12, 4 + int(obj_share // 15) * 2 + int(total_abs // 200_000)), f"对象占本人资金往来{obj_share:.1f}%，累计{_fmt_money_human(total_abs)}")
    cash_combo = int(cash_cnt > 0 and quick_cnt > 0) + int(cash_cnt > 0 and ratio >= 0.40) + int(cash_cnt > 0 and net_ratio <= 0.30 and in_sum > 0 and out_sum > 0)
    if cash_combo:
        add("疑似资金代持/白手套线索", "WG-12", "现金叠加快进快出/对倒", min(12, cash_combo * 4 + min(4, cash_in_cnt + cash_out_cnt)), f"现金{cash_cnt}笔，叠加快进快出/对倒/净额小流水大特征{cash_combo}类")

    # 三、疑似利益输送线索：重视职务/机构对象向本人转入，叠加敏感时点、通信、现金和项目业务文本。
    if has_title and in_sum > 0 and (max_in_amt >= 3_000 or in_sum >= 10_000):
        score = min(24, 4 + in_cnt * 2 + int(in_sum // 50_000) * 3 + int(max_in_amt // 50_000) * 2)
        add("疑似利益输送线索", "INT-01", "职务对象转入", score, f"对象职务：{safe_str(row.get('对象职务', '')).strip()}；转入{in_cnt}笔{_fmt_money_human(in_sum)}，最大转入{_fmt_money_human(max_in_amt)}")
    if is_company and in_sum > 0 and (max_in_amt >= 5_000 or in_sum >= 30_000):
        score = min(24, 4 + in_cnt * 2 + int(in_sum // 100_000) * 4 + int(max_in_amt // 100_000) * 2)
        add("疑似利益输送线索", "INT-02", "机构类对象转入", score, f"机构类对象转入{in_cnt}笔{_fmt_money_human(in_sum)}，最大转入{_fmt_money_human(max_in_amt)}")
    if benefit_kw > 0:
        add("疑似利益输送线索", "INT-03", "利益输送文本", min(25, benefit_kw * 5), f"利益输送类关键词{benefit_kw}笔")
    sensitive_in = _safe_int_num(row.get("敏感时点转入次数", 0))
    if in_sum >= 3_000 and sensitive_in > 0 and (has_title or is_company):
        add("疑似利益输送线索", "INT-04", "敏感时点转入", min(18, sensitive_in * 4 + int(in_sum // 100_000)), f"敏感时点转入{sensitive_in}笔，转入{_fmt_money_human(in_sum)}")
    if split_cnt >= 3 and split_amt >= 30_000:
        add("疑似利益输送线索", "INT-05", "拆分转入", min(18, split_cnt * 3), f"5千至5万元区间转入{split_cnt}笔，累计{_fmt_money_human(split_amt)}")
    if in_months >= 2 and in_cnt >= 3:
        add("疑似利益输送线索", "INT-06", "长期稳定转入", min(16, (in_months - 1) * 3 + in_cnt), f"跨{in_months}个月转入，转入{in_cnt}笔")
    if in_sum > 0 and (comm_cnt >= 5 or comm_night >= 1 or comm_fest > 0) and (has_title or is_company):
        score = min(16, max(1, comm_cnt // 5) * 2 + comm_night * 3 + comm_fest * 3)
        add("疑似利益输送线索", "INT-07", "资金叠加通信", score, f"存在转入资金，同时通信{comm_cnt}次、深夜通信{comm_night}次、节日通信{comm_fest}次")
    if (has_title or is_company) and cash_in_cnt > 0:
        add("疑似利益输送线索", "INT-08", "职务/机构对象现金转入", min(16, cash_in_cnt * 4 + int(in_sum // 100_000)), f"现金/ATM/卡存类转入{cash_in_cnt}笔，转入{_fmt_money_human(in_sum)}")
    if (has_title or is_company) and round_in_cnt >= 2 and in_sum >= 20_000:
        add("疑似利益输送线索", "INT-09", "整额转入", min(12, round_in_cnt * 2 + int(in_sum // 100_000)), f"整额大额转入{round_in_cnt}笔，转入{_fmt_money_human(in_sum)}")
    if (has_title or is_company) and in_ratio_pct >= 70 and in_sum >= 20_000:
        add("疑似利益输送线索", "INT-10", "高占比转入", min(12, 4 + int(in_ratio_pct // 20) * 2 + int(in_sum // 100_000)), f"转入占比{_fmt_pct(in_ratio_pct)}，转入{_fmt_money_human(in_sum)}")
    if is_company and in_sum > 0 and project_kw > 0:
        add("疑似利益输送线索", "INT-11", "项目/工程/服务类机构转入", min(14, project_kw * 3 + int(in_sum // 100_000) * 2), f"项目/工程/服务类关键词{project_kw}笔，机构转入{_fmt_money_human(in_sum)}")

    return rules


def _score_profile(profile: pd.DataFrame) -> pd.DataFrame:
    if profile is None or profile.empty:
        return pd.DataFrame()
    rows: List[Dict[str, Any]] = []
    for _, r in profile.iterrows():
        score_pack = _score_relation_row(r)
        base = r.to_dict()
        for clue_type in CLUE_TYPES:
            rules = score_pack.get(clue_type, [])
            score = min(100, sum(int(x.get("分值", 0)) for x in rules))
            base[f"{clue_type}-风险分值"] = score
            base[f"{clue_type}-风险等级"] = _risk_grade(score)
            base[f"{clue_type}-规则编号"] = "、".join(dict.fromkeys([safe_str(x.get("规则编号", "")).strip() for x in rules if safe_str(x.get("规则编号", "")).strip()]))
            base[f"{clue_type}-命中规则"] = "；".join(f"{x.get('规则编号')} {x.get('规则名称')}（+{x.get('分值')}）" for x in rules)
            base[f"{clue_type}-证据摘要"] = "；".join(safe_str(x.get("证据", "")).strip() for x in rules if safe_str(x.get("证据", "")).strip())
        rows.append(base)
    out = pd.DataFrame(rows).replace({np.nan: ""})
    score_cols = [f"{t}-风险分值" for t in CLUE_TYPES]
    out["最高风险分值"] = out[score_cols].apply(lambda s: max(pd.to_numeric(s, errors="coerce").fillna(0)), axis=1)
    out.sort_values(["最高风险分值", "资金往来总额", "通信次数"], ascending=[False, False, False], inplace=True, kind="mergesort")
    return out.reset_index(drop=True)


def _empty_clue_row(person: str, clue_type: str) -> Dict[str, Any]:
    return {
        "查询对象": person,
        "线索类型": clue_type,
        "判断结论": "未命中",
        "风险等级": "无",
        "风险分值": 0,
        "关系对象": "",
        "交易对方姓名": "",
        "对方职务": "",
        "一句话提示": f"未命中｜{clue_type}｜未达到自动提示阈值",
        "规则编号": "",
        "命中规则": "",
        "判断逻辑": "未达到该类别自动筛查阈值。低分或单因素异常暂不进入自动提示。",
        "关键证据": "",
        "核心证据": "",
        "证据样例": "",
        "证据笔数": 0,
        "涉及金额": 0,
        "最大单笔": 0,
        "首次交易时间": "",
        "末次交易时间": "",
        "核查建议": "",
        "建议核查方向": "",
        "证据明细文件": "对应线索sheet内的“证据样例/关键证据”字段",
    }


def _advice_for_clue_type(clue_type: str) -> str:
    if clue_type == "疑似情感关系线索":
        return "结合双方身份关系、通信内容、节日前后接触、同住同机及资金用途核实，避免仅凭单笔特殊金额定性。"
    if clue_type == "疑似资金代持/白手套线索":
        return "核查资金最终来源和去向、账户实际控制人、是否替他人收付款、是否存在快进快出或资产代持安排。"
    if clue_type == "疑似利益输送线索":
        return "核查对方身份、职务或经营主体与查询对象职权事项、项目审批、工程采购、管理服务关系及资金真实性质。"
    return "结合原始凭证和外部背景进一步核查。"


def _evidence_examples_from_df(g: pd.DataFrame, limit: int = 5) -> str:
    if g is None or g.empty:
        return ""
    tmp = g.copy()
    tmp["__abs_sort__"] = pd.to_numeric(tmp.get("__abs_amt__", tmp.get("交易金额", 0)), errors="coerce").abs()
    tmp["__dt_sort__"] = pd.to_datetime(tmp.get("交易时间", ""), errors="coerce")
    tmp = tmp.sort_values(["__abs_sort__", "__dt_sort__"], ascending=[False, True], kind="mergesort").head(limit)
    pieces: List[str] = []
    for _, r in tmp.iterrows():
        dt = r.get("__dt_sort__", pd.NaT)
        date_s = dt.strftime("%Y-%m-%d") if pd.notna(dt) else safe_str(r.get("交易时间", "")).strip() or "日期缺失"
        direction = safe_str(r.get("__方向说明__", "")).strip() or safe_str(r.get("借贷标志", "")).strip() or "方向未明"
        amt = _fmt_money_human(r.get("__abs_sort__", 0))
        text = " ".join(safe_str(r.get(c, "")).strip() for c in ["交易摘要", "摘要", "备注", "注释", "备注/注释", "商户名称"])
        text = re.sub(r"\s+", " ", text).strip()
        if len(text) > 30:
            text = text[:30] + "…"
        pieces.append(f"{date_s} {direction} {amt}" + (f"（{text}）" if text else ""))
    return "；".join(pieces)


def _select_evidence_rows(d: pd.DataFrame, obj: str, clue_type: str, profile_row: pd.Series, max_rows: int = 80) -> pd.DataFrame:
    if d is None or d.empty:
        return pd.DataFrame()
    g = d[d["关系对象"].map(safe_str).str.strip() == safe_str(obj).strip()].copy()
    if g.empty:
        return pd.DataFrame()
    if clue_type == "疑似情感关系线索":
        mask = g["__special_emotion_amount__"] | g["__emotion_keyword__"] | g["__lifecare_keyword__"] | g["__festival__"].isin(["七夕节", "5月20日"]) | g["__late_night__"] | g["__offwork__"] | g["__small_amt__"]
    elif clue_type == "疑似资金代持/白手套线索":
        quick_idx = set()
        for x in safe_str(profile_row.get("快进快出流水索引", "")).split("、"):
            x = x.strip()
            if x.isdigit():
                quick_idx.add(int(x))
        mask = g["__holding_keyword__"] | g["__cash_signal__"] | g["__mid_amt__"] | (g["__abs_amt__"] >= 10_000) | g.index.isin(quick_idx)
    else:
        mask = g["__is_in__"] | g["__benefit_keyword__"] | ((g["__abs_amt__"] >= 5_000) & (g["__offwork__"] | g["__late_night__"] | g["节假日"].map(safe_str).str.strip().eq("节假日")))
    ev = g[mask].copy()
    if ev.empty:
        ev = g.copy()
    ev["__abs_sort__"] = ev["__abs_amt__"].abs()
    ev = ev.sort_values(["__abs_sort__", "交易时间"], ascending=[False, True], kind="mergesort").head(max_rows)
    keep = [
        "查询对象", "关系对象", "交易对方姓名", "对方职务", "借贷标志", "交易金额", "交易时间", "__方向说明__", "__festival__",
        "节假日", "星期", "交易摘要", "摘要", "备注", "注释", "备注/注释", "商户名称", "交易类型", "反馈单位",
        "查询账户", "查询卡号", "交易对方账户", "交易对方卡号", "交易对方账号开户行", "交易流水号", "来源文件",
    ]
    for c in keep:
        if c not in ev.columns:
            ev[c] = ""
    ev = ev[keep].rename(columns={"__方向说明__": "资金方向", "__festival__": "特殊节日", "对方职务": "交易对方职务"})
    ev["交易时间"] = pd.to_datetime(ev["交易时间"], errors="coerce").dt.strftime("%Y-%m-%d %H:%M:%S").fillna("")
    return ev.reset_index(drop=True)


def _shorten_hit_rules(rule_text: Any, limit: int = 6) -> str:
    parts = [p.strip() for p in safe_str(rule_text).split("；") if p.strip()]
    if len(parts) <= limit:
        return "；".join(parts)
    return "；".join(parts[:limit]) + "；…"


def _compact_rule_names(rule_text: Any, limit: int = 4) -> str:
    parts = [p.strip() for p in safe_str(rule_text).split("；") if p.strip()]
    out: List[str] = []
    for p in parts[:limit]:
        m = re.match(r"([A-Z]+-\d+)\s*(.+?)(?:（|$)", p)
        if m:
            out.append(f"{m.group(1)}{m.group(2).strip()}")
        else:
            out.append(p.split("（", 1)[0])
    if len(parts) > limit:
        out.append("…")
    return "、".join([x for x in out if x])


def _clue_one_line(person: str, obj: str, clue_type: str, score: Any, row: pd.Series, rules: Any) -> str:
    risk = _risk_grade(score)
    obj_s = safe_str(obj).strip() or "关系对象缺失"
    rule_s = _compact_rule_names(rules, limit=4)
    ev = _profile_metrics_sentence(row)
    if len(ev) > 90:
        ev = ev[:90] + "…"
    type_s = safe_str(clue_type).replace("疑似", "").replace("线索", "")
    return f"{risk}风险｜{type_s}｜{obj_s}｜{rule_s}" + (f"｜{ev}" if ev else "")


def _concise_clue_view(summary: pd.DataFrame) -> pd.DataFrame:
    if summary is None or summary.empty:
        return pd.DataFrame()
    cols = [
        "判断结论", "线索类型", "风险等级", "风险分值", "关系对象", "对方职务",
        "一句话提示", "命中规则", "关键证据", "证据样例", "证据笔数", "涉及金额", "首次交易时间", "末次交易时间", "核查建议",
    ]
    d = summary.copy()
    if "一句话提示" not in d.columns:
        d["一句话提示"] = ""
    if "关键证据" not in d.columns and "核心证据" in d.columns:
        d["关键证据"] = d["核心证据"]
    if "核查建议" not in d.columns and "建议核查方向" in d.columns:
        d["核查建议"] = d["建议核查方向"]
    for c in cols:
        if c not in d.columns:
            d[c] = ""
    return d[cols].replace({np.nan: ""})


def _build_summary_and_detail(scored_profile: pd.DataFrame, prepared_txn: pd.DataFrame, person: str) -> Tuple[pd.DataFrame, pd.DataFrame]:
    summary_rows: List[Dict[str, Any]] = []
    detail_frames: List[pd.DataFrame] = []
    hit_type_counter = {t: 0 for t in CLUE_TYPES}

    if scored_profile is None or scored_profile.empty:
        return pd.DataFrame([_empty_clue_row(person, t) for t in CLUE_TYPES]), pd.DataFrame()

    for _, r in scored_profile.iterrows():
        obj = safe_str(r.get("关系对象", "")).strip()
        if not obj:
            continue
        for clue_type in CLUE_TYPES:
            score = _safe_int_num(r.get(f"{clue_type}-风险分值", 0))
            if not _is_hit_score(score):
                continue
            hit_type_counter[clue_type] += 1
            rule_ids = safe_str(r.get(f"{clue_type}-规则编号", "")).strip()
            rules = _shorten_hit_rules(r.get(f"{clue_type}-命中规则", ""), limit=5)
            evidence_logic = safe_str(r.get(f"{clue_type}-证据摘要", "")).strip()
            key_evidence = _profile_metrics_sentence(r)
            one_line = _clue_one_line(person, obj, clue_type, score, r, rules)
            ev = _select_evidence_rows(prepared_txn, obj, clue_type, r)
            if not ev.empty:
                ev.insert(0, "线索类型", clue_type)
                ev.insert(1, "风险分值", score)
                ev.insert(2, "风险等级", _risk_grade(score))
                ev.insert(3, "规则编号", rule_ids)
                ev.insert(4, "命中规则", rules)
                ev.insert(5, "一句话提示", one_line)
                detail_frames.append(ev)
                evidence_cnt = len(ev)
                first_dt = pd.to_datetime(ev.get("交易时间", ""), errors="coerce").min()
                last_dt = pd.to_datetime(ev.get("交易时间", ""), errors="coerce").max()
                sample = _evidence_examples_from_df(prepared_txn[prepared_txn["关系对象"].map(safe_str).str.strip() == obj], limit=3)
            else:
                synthetic = pd.DataFrame([{
                    "线索类型": clue_type,
                    "风险分值": score,
                    "风险等级": _risk_grade(score),
                    "规则编号": rule_ids,
                    "命中规则": rules,
                    "一句话提示": one_line,
                    "查询对象": person,
                    "关系对象": obj,
                    "交易对方姓名": safe_str(r.get("交易对方姓名", "")).strip(),
                    "交易对方职务": safe_str(r.get("对象职务", "")).strip(),
                    "证据来源": "关系对象画像（通信/同住/同机或统计特征）",
                    "证据摘要": key_evidence,
                }])
                detail_frames.append(synthetic)
                evidence_cnt = _safe_int_num(r.get("通信次数", 0)) + _safe_int_num(r.get("旅馆同住次数", 0)) + _safe_int_num(r.get("同机次数", 0))
                first_dt = pd.NaT
                last_dt = pd.NaT
                sample = key_evidence

            advice = _advice_for_clue_type(clue_type)
            summary_rows.append({
                "判断结论": "命中",
                "线索类型": clue_type,
                "风险等级": _risk_grade(score),
                "风险分值": score,
                "关系对象": obj,
                "对方职务": safe_str(r.get("对象职务", "")).strip(),
                "一句话提示": one_line,
                "命中规则": rules,
                "关键证据": key_evidence,
                "证据样例": sample,
                "证据笔数": int(evidence_cnt),
                "涉及金额": float(_safe_float_num(r.get("资金往来总额", 0))),
                "最大单笔": float(_safe_float_num(r.get("最大单笔", 0))),
                "首次交易时间": safe_str(r.get("首次交易时间", "")).strip() or (first_dt.strftime("%Y-%m-%d %H:%M:%S") if pd.notna(first_dt) else ""),
                "末次交易时间": safe_str(r.get("末次交易时间", "")).strip() or (last_dt.strftime("%Y-%m-%d %H:%M:%S") if pd.notna(last_dt) else ""),
                "核查建议": advice,
                # 兼容原汇总报告和旧逻辑字段。
                "查询对象": person,
                "交易对方姓名": safe_str(r.get("交易对方姓名", "")).strip() or obj,
                "规则编号": rule_ids,
                "判断逻辑": evidence_logic,
                "核心证据": key_evidence,
                "建议核查方向": advice,
                "证据明细文件": "对应线索sheet内的“证据样例/关键证据”字段",
            })

    for clue_type, cnt in hit_type_counter.items():
        if cnt == 0:
            empty = _empty_clue_row(person, clue_type)
            empty["一句话提示"] = f"未命中｜{clue_type}｜未达到自动提示阈值"
            empty["关键证据"] = ""
            empty["核心证据"] = ""
            empty["核查建议"] = empty.get("建议核查方向", "")
            summary_rows.append(empty)

    cols = [
        "判断结论", "线索类型", "风险等级", "风险分值", "关系对象", "对方职务", "一句话提示", "命中规则",
        "关键证据", "证据样例", "证据笔数", "涉及金额", "最大单笔", "首次交易时间", "末次交易时间", "核查建议",
        "查询对象", "交易对方姓名", "规则编号", "判断逻辑", "核心证据", "建议核查方向", "证据明细文件",
    ]
    summary = pd.DataFrame(summary_rows)
    for c in cols:
        if c not in summary.columns:
            summary[c] = ""
    summary = summary[cols].replace({np.nan: ""})
    summary["__hit_rank__"] = summary["判断结论"].map({"命中": 1, "未命中": 0}).fillna(0)
    summary["__risk_rank__"] = summary["风险等级"].map(_risk_rank)
    summary["风险分值"] = pd.to_numeric(summary["风险分值"], errors="coerce").fillna(0)
    summary["涉及金额"] = pd.to_numeric(summary["涉及金额"], errors="coerce").fillna(0)
    summary["证据笔数"] = pd.to_numeric(summary["证据笔数"], errors="coerce").fillna(0)
    summary.sort_values(["__hit_rank__", "__risk_rank__", "风险分值", "涉及金额", "证据笔数"], ascending=[False, False, False, False, False], inplace=True, kind="mergesort")
    summary = summary.drop(columns=["__hit_rank__", "__risk_rank__"], errors="ignore").reset_index(drop=True)
    detail = pd.concat(detail_frames, ignore_index=True).replace({np.nan: ""}) if detail_frames else pd.DataFrame()
    return summary, detail


def _clue_type_sheet_name(clue_type: str) -> str:
    mp = {
        "疑似情感关系线索": "疑似情感关系",
        "疑似资金代持/白手套线索": "疑似资金代持白手套",
        "疑似利益输送线索": "疑似利益输送",
    }
    return mp.get(clue_type, safe_str(clue_type).replace("/", ""))[:31]


def _three_clue_sheets(summary: pd.DataFrame) -> Dict[str, pd.DataFrame]:
    """生成只含三类线索 sheet 的导出结构。

    每个 sheet 对应一种线索类型，只放简洁判断结果，避免再输出规则说明、证据明细、画像表等辅助 sheet。
    """
    sheets: Dict[str, pd.DataFrame] = {}
    if summary is None or summary.empty:
        for clue_type in CLUE_TYPES:
            sheets[_clue_type_sheet_name(clue_type)] = pd.DataFrame()
        return sheets

    d = summary.copy().replace({np.nan: ""})
    if "线索类型" not in d.columns:
        d["线索类型"] = ""

    for clue_type in CLUE_TYPES:
        sub = d[d["线索类型"].map(safe_str).str.strip() == clue_type].copy()
        sheets[_clue_type_sheet_name(clue_type)] = _concise_clue_view(sub)
    return sheets


def generate_suspicious_clues(df_person: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """按“关系对象画像 + 分值研判”生成三类疑似线索，Excel 只输出三类线索 sheet。"""
    if df_person is None or df_person.empty:
        return pd.DataFrame(), pd.DataFrame()

    person = safe_str(df_person["查询对象"].iat[0] if "查询对象" in df_person.columns else "").strip() or "未知"
    prepared = _prepare_clue_df(df_person, person)
    bank_profile = build_bank_counterparty_profile(df_person, person)
    comm_profile = build_comm_counterparty_profile(person)
    co_travel_profile = build_co_travel_profile(person)
    profile = merge_relation_profiles(bank_profile, comm_profile, co_travel_profile, person)
    scored_profile = _score_profile(profile)
    summary, detail = _build_summary_and_detail(scored_profile, prepared, person)

    # 按用户要求：生成的 Excel 只保留 3 个 sheet，分别对应三类线索。
    sheets = _three_clue_sheets(summary)

    out_file = out_report_dir(person) / f"{person}-关系风险画像及三类疑似线索"
    save_excel_sheets_auto_width(sheets, out_file, index=False)
    STATE.CLUE_SUMMARY_BY_PERSON[person] = summary
    STATE.CLUE_DETAIL_BY_PERSON[person] = detail
    print(f"✅ 已生成三类疑似线索：{out_file.with_suffix('.xlsx')}（仅含3个sheet）")
    return summary, detail


def export_all_clue_summary() -> None:
    frames = [v for v in STATE.CLUE_SUMMARY_BY_PERSON.values() if v is not None and not v.empty]
    if not frames:
        return
    all_summary = pd.concat(frames, ignore_index=True).replace({np.nan: ""})
    # 汇总工作簿同样只保留 3 个 sheet，分别对应三类线索。
    sheets = _three_clue_sheets(all_summary)
    save_excel_sheets_auto_width(sheets, out_root() / "所有人-关系风险画像及三类疑似线索汇总", index=False)
    print("✅ 已生成全部人员三类疑似线索汇总：分析结果/所有人-关系风险画像及三类疑似线索汇总.xlsx（仅含3个sheet）")

# ========= 汇总报告（txt） =========
def _fmt_human_num(n: Any) -> str:
    try:
        if n is None or (isinstance(n, float) and np.isnan(n)):
            return "0"
        x = float(n)
    except Exception:
        s = safe_str(n).strip()
        return s if s else "0"
    ax = abs(x)
    def _strip(v: float) -> str: return f"{v:.2f}".rstrip("0").rstrip(".")
    if ax >= 1e8: return _strip(x / 1e8) + "亿"
    if ax >= 1e4: return _strip(x / 1e4) + "万"
    if ax >= 1e3: return _strip(x / 1e3) + "千"
    return _strip(x)

def _fmt_money_human(x: Any) -> str:
    try:
        if x is None or (isinstance(x, float) and np.isnan(x)):
            return "0元"
        v = float(x)
    except Exception:
        s = safe_str(x).strip()
        return (s + "元") if s else "0元"
    av = abs(v)
    def _strip(vv: float) -> str: return f"{vv:.2f}".rstrip("0").rstrip(".")
    if av >= 1e8: return _strip(v / 1e8) + "亿元"
    if av >= 1e4: return _strip(v / 1e4) + "万元"
    return _strip(v) + "元"

def _build_top8_partner_narratives(d: pd.DataFrame, person: str, top_n: int = 8) -> List[str]:
    if d is None or d.empty:
        return []
    person = safe_str(person).strip() or "未知"
    dd = d.copy()
    for c in ["交易对方姓名","交易时间","交易金额","借贷标志"]:
        if c not in dd.columns:
            dd[c] = ""
    dd["交易时间"] = pd.to_datetime(dd["交易时间"], errors="coerce")
    dd["__amt__"] = pd.to_numeric(dd["交易金额"], errors="coerce")
    dd["__abs_amt__"] = dd["__amt__"].abs()
    dd["交易对方姓名"] = dd["交易对方姓名"].map(safe_str).str.strip()
    dd = dd[(dd["交易对方姓名"] != "") & (dd["交易对方姓名"] != " ")]
    if dd.empty:
        return []
    dd = dd[dd["交易对方姓名"].map(lambda x: bool(CH_NAME_2_4_RE.fullmatch(x)))]
    if dd.empty:
        return []

    is_in = dd["借贷标志"].map(safe_str).str.strip().eq("进")
    is_out = dd["借贷标志"].map(safe_str).str.strip().eq("出")
    fb_in = (~is_in) & (~is_out) & (dd["__amt__"] > 0)
    fb_out = (~is_in) & (~is_out) & (dd["__amt__"] < 0)
    dd["__is_in__"] = is_in | fb_in
    dd["__is_out__"] = is_out | fb_out

    grp = dd.groupby("交易对方姓名", dropna=False)["__abs_amt__"].sum().sort_values(ascending=False)
    top_names = grp.head(top_n).index.tolist()
    if not top_names:
        return []

    out_lines: List[str] = [f"（交易对手分析：仅统计2-4个汉字姓名，按交易金额最大的前{len(top_names)}名摘要）"]
    for i, name in enumerate(top_names, start=1):
        g = dd[dd["交易对方姓名"] == name].sort_values("交易时间", kind="mergesort")
        if g.empty:
            continue
        g_pick = g[g["__abs_amt__"] >= 10_000].copy()
        if g_pick.empty:
            g_pick = g.nlargest(5, "__abs_amt__").sort_values("交易时间", kind="mergesort")
        if len(g_pick) > 12:
            g_pick = g_pick.nlargest(12, "__abs_amt__").sort_values("交易时间", kind="mergesort")

        pieces = []
        for _, r in g_pick.iterrows():
            dt = r.get("交易时间", pd.NaT)
            date_s = dt.strftime("%Y-%m-%d") if pd.notna(dt) else "日期缺失"
            amt_h = _fmt_money_human(r.get("__abs_amt__", np.nan))
            if bool(r.get("__is_in__", False)):
                pieces.append(f"{date_s}，{name}向{person}转账{amt_h}")
            elif bool(r.get("__is_out__", False)):
                pieces.append(f"{date_s}，{person}向{name}转账{amt_h}")
            else:
                pieces.append(f"{date_s}，{person}与{name}转账{amt_h}")

        in_cnt = int(g["__is_in__"].sum())
        out_cnt = int(g["__is_out__"].sum())
        in_sum = _safe_series_sum(g.loc[g["__is_in__"], "__abs_amt__"])
        out_sum = _safe_series_sum(g.loc[g["__is_out__"], "__abs_amt__"])

        narrative = "；".join(pieces) + "。"
        tail = f"合计：{name}向{person}转账{_fmt_human_num(in_cnt)}笔{_fmt_money_human(in_sum)}，{person}向{name}转账{_fmt_human_num(out_cnt)}笔{_fmt_money_human(out_sum)}。"
        out_lines.append(f"{i}. {narrative}{tail}")
    return out_lines

def build_person_report_txt(root: Path, person: str, df_person: pd.DataFrame) -> Path:
    person = safe_str(person).strip() or "未知"
    now = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    lines: List[str] = []
    lines += [f"{person}交易流水、通信记录及不动产信息综合情况报告", "（自动生成）", "", f"生成时间：{now}", f"工作目录：{root}", ""]
    lines.append("一、交易流水情况")

    if df_person is None or df_person.empty:
        lines += ["经检索整理，未发现可用于分析的交易流水数据。", ""]
    else:
        d = df_person.copy()
        need_cols = ["借贷标志","交易金额","交易时间","反馈单位","交易对方姓名","对方职务","账户余额","现金标志","交易类型","查询账户","查询卡号",
                     "交易对方账号开户行","交易对方账户","交易对方卡号","交易摘要","备注","注释"]
        for c in need_cols:
            if c not in d.columns:
                d[c] = ""

        d["交易时间"] = pd.to_datetime(d["交易时间"], errors="coerce")
        amt_raw = pd.to_numeric(d["交易金额"], errors="coerce")
        d["__amt__"] = amt_raw
        d["__abs_amt__"] = amt_raw.abs()

        is_in = d["借贷标志"].map(safe_str).str.strip().eq("进")
        is_out = d["借贷标志"].map(safe_str).str.strip().eq("出")
        is_in = is_in | ((~is_in) & (~is_out) & (d["__amt__"] > 0))
        is_out = is_out | ((~is_in) & (~is_out) & (d["__amt__"] < 0))

        total_cnt = len(d)
        total_amt = d["__abs_amt__"].sum(skipna=True)
        in_amt = d["__abs_amt__"].where(is_in, 0).sum(skipna=True)
        out_amt = d["__abs_amt__"].where(is_out, 0).sum(skipna=True)
        net_in = in_amt - out_amt

        tmin = d["交易时间"].min() if d["交易时间"].notna().any() else pd.NaT
        tmax = d["交易时间"].max() if d["交易时间"].notna().any() else pd.NaT
        max_in = d["__abs_amt__"].where(is_in, np.nan).max(skipna=True)
        max_out = d["__abs_amt__"].where(is_out, np.nan).max(skipna=True)

        cash = d[_cash_signal_mask(d) & (d["__abs_amt__"] >= 10_000)]
        big = d[d["__abs_amt__"] >= 500_000]

        if pd.notna(tmin) and pd.notna(tmax):
            lines.append(
                f"经对{person}相关银行交易流水进行归集与规范化处理，时间跨度为{tmin.strftime('%Y-%m-%d %H:%M:%S')}至{tmax.strftime('%Y-%m-%d %H:%M:%S')}。"
                f"期间共发生交易{_fmt_human_num(total_cnt)}笔，涉及金额合计约{_fmt_money_human(total_amt)}；"
                f"其中资金流入约{_fmt_money_human(in_amt)}、资金流出约{_fmt_money_human(out_amt)}，净流入约{_fmt_money_human(net_in)}。"
            )
        else:
            lines.append(
                f"经对{person}相关银行交易流水进行归集与规范化处理，共发生交易{_fmt_human_num(total_cnt)}笔，涉及金额合计约{_fmt_money_human(total_amt)}；"
                f"其中资金流入约{_fmt_money_human(in_amt)}、资金流出约{_fmt_money_human(out_amt)}，净流入约{_fmt_money_human(net_in)}。"
                "但交易时间字段存在缺失或格式异常，时间范围无法准确提取。"
            )

        lines.append(
            f"从单笔特征看，单笔最大流入约{_fmt_money_human(max_in)}，单笔最大流出约{_fmt_money_human(max_out)}；"
            f"从敏感特征看，存取现（单笔不低于1万元）{_fmt_human_num(len(cash))}笔，大额资金交易（单笔不低于50万元）{_fmt_human_num(len(big))}笔。"
        )

        top8_lines = _build_top8_partner_narratives(d, person=person, top_n=8)
        if top8_lines:
            lines.extend(top8_lines)

        # 余额汇总（银行余额总额）
        bal = d.copy()
        bal["交易时间"] = pd.to_datetime(bal["交易时间"], errors="coerce")
        bal["账户余额"] = pd.to_numeric(bal["账户余额"], errors="coerce")
        bal["反馈单位"] = bal["反馈单位"].map(safe_str).replace("", "（未知银行）")
        acct_id = bal["查询账户"].map(safe_str).str.strip()
        acct_id = acct_id.where(acct_id != "", bal["查询卡号"].map(safe_str).str.strip())
        bal["__acct_id__"] = acct_id.replace("", "（账户未知）")

        bal_valid = bal.dropna(subset=["账户余额"]).copy()
        if not bal_valid.empty:
            if bal_valid["交易时间"].notna().any():
                bal_valid = bal_valid.sort_values("交易时间", kind="mergesort")
            last_per_acct = bal_valid.groupby(["反馈单位","__acct_id__"], dropna=False).tail(1)[["反馈单位","__acct_id__","账户余额","交易时间"]]
            bank_sum = last_per_acct.groupby("反馈单位", dropna=False)["账户余额"].sum().reset_index().sort_values("账户余额", ascending=False, kind="mergesort")
            total_balance = float(bank_sum["账户余额"].sum()) if "账户余额" in bank_sum.columns else 0.0
            parts = [f"{safe_str(r['反馈单位']).strip() or '（未知银行）'}合计约{_fmt_money_human(_safe_float_num(r.get('账户余额', 0)))}" for _, r in bank_sum.iterrows()]
            lines.append("从余额字段可提取的情况看，按各银行账户末笔余额汇总后，" + "；".join(parts) + f"。银行余额总额约{_fmt_money_human(total_balance)}。")
        else:
            lines.append("余额情况：账户余额字段缺失或无法解析，未能形成各银行余额汇总及总额。")

        lines.append("")

    # 二、通信记录（按人名精确匹配）
    lines.append("二、通信记录情况")
    st = STATE.COMM_STATS_BY_PERSON.get(person)
    rec_cnt = _safe_int_num(STATE.COMM_RECORDS_BY_PERSON.get(person, 0))

    if st is None or st.empty:
        lines += [f"经检索，本次未能获取“{person}”对应的已生成通信统计表（请确认是否存在“*-{COMM_SUFFIX}.xlsx/.xls”且其内记录包含该人员）。", ""]
    else:
        lines.append(
            f"经对已生成的通信统计结果进行汇总整理，按号码归并后涉及{_fmt_human_num(len(st))}个对方号码；"
            + (f"累计通信记录约{_fmt_human_num(rec_cnt)}条。" if rec_cnt else "")
            + "其中非工作时间、深夜时段及通话时长较长的通信行为，建议结合对象关系与事件背景进一步核验。"
        )
        topc = st.copy()
        if "通信次数" in topc.columns and "通话≥3分钟次数" in topc.columns:
            topc = topc.sort_values(["通信次数","通话≥3分钟次数"], ascending=[False, False], kind="mergesort")
        elif "通信次数" in topc.columns:
            topc = topc.sort_values(["通信次数"], ascending=[False], kind="mergesort")
        topc = topc.head(5)

        desc_parts = []
        for _, r in topc.iterrows():
            phone = safe_str(r.get("对方号码","")).strip() or "（号码缺失）"
            nm = safe_str(r.get("姓名","")).strip()
            tt = safe_str(r.get("职务","")).strip()
            c1 = _safe_int_num(r.get("通信次数", 0))
            c2 = _safe_int_num(r.get("非工作时间通信次数", 0))
            c3 = _safe_int_num(r.get("深夜通信次数(23–5)", 0))
            c4 = _safe_int_num(r.get("通话≥3分钟次数", 0))
            who = phone + (f"（{nm}/{tt}）" if (nm and tt) else (f"（{nm}）" if nm else ""))
            desc_parts.append(f"{who}：通信{_fmt_human_num(c1)}次，非工作时间{_fmt_human_num(c2)}次、深夜{_fmt_human_num(c3)}次、通话≥3分钟{_fmt_human_num(c4)}次")
        if desc_parts:
            lines.append("从高频对象看，通信较为频繁的对方号码主要集中在以下对象：" + "；".join(desc_parts) + "。")
        lines.append("")

    # 三、不动产
    lines.append("三、不动产信息情况")
    re_person, re_file_cnt = _read_realestate_for_person_by_filename(root, person)
    if not re_person:
        if re_file_cnt == 0:
            lines.append(f"经检索，未发现文件名以“{person}-全省不动产”开头的不动产数据文件，故本次未形成可用的不动产摘要。")
        else:
            lines.append(f"已检索到与“{person}-全省不动产”匹配的不动产文件{_fmt_human_num(re_file_cnt)}个，但可解析的有效记录为空或关键字段缺失。")
        lines.append("")
    else:
        lines.append(f"经检索汇总，已匹配不动产文件{_fmt_human_num(re_file_cnt)}个，提取有效记录{_fmt_human_num(len(re_person))}条。现按字段摘述房屋坐落、建筑面积、交易价格（万元）、登记时间等要素如下。")
        take = re_person[:3]
        sents = []
        for r in take:
            loc = r.get("房屋坐落","") or "（坐落未填写）"
            area = r.get("建筑面积","") or "（建筑面积未填写）"
            price = r.get("交易价格（万元）","") or "（交易价格未填写）"
            regt = r.get("登记时间","") or "（登记时间未填写）"
            sents.append(f"房屋坐落：{loc}；建筑面积：{area}；交易价格（万元）：{price}；登记时间：{regt}")
        lines.append("摘述要点：" + "；".join(sents) + "。")
        lines.append("")

    # 四、三类线索自动研判
    lines.append("四、关系风险线索自动研判情况")
    clue_summary = STATE.CLUE_SUMMARY_BY_PERSON.get(person)
    if clue_summary is None or clue_summary.empty:
        lines.append("本次尚未生成关系风险画像及三类疑似线索判断表。")
    else:
        cs = clue_summary.copy()
        if "判断结论" in cs.columns:
            hit_cs = cs[cs["判断结论"].map(safe_str).str.strip() == "命中"].copy()
        else:
            hit_cs = cs[cs.get("证据笔数", pd.Series([0] * len(cs))).map(lambda x: _safe_int_num(x) > 0)].copy()
        if hit_cs.empty:
            lines.append("经对象画像和分值研判，未达到“疑似情感关系线索、疑似资金代持/白手套线索、疑似利益输送线索”的自动提示阈值。")
        else:
            type_counts = hit_cs.groupby("线索类型", dropna=False)["交易对方姓名"].count().to_dict() if "线索类型" in hit_cs.columns else {}
            parts = [f"{k}{_fmt_human_num(v)}项" for k, v in type_counts.items()]
            lines.append("经对象画像和分值研判，命中" + "、".join(parts) + "。具体判断已按三类线索分别导出至个人报告目录下的“关系风险画像及三类疑似线索.xlsx”。")
            tmp = hit_cs.copy()
            tmp["__risk_rank__"] = tmp["风险等级"].map(_risk_rank) if "风险等级" in tmp.columns else 0
            tmp["风险分值"] = pd.to_numeric(tmp.get("风险分值", 0), errors="coerce").fillna(0)
            tmp["涉及金额"] = pd.to_numeric(tmp.get("涉及金额", 0), errors="coerce").fillna(0)
            tmp["证据笔数"] = pd.to_numeric(tmp.get("证据笔数", 0), errors="coerce").fillna(0)
            top_cs = tmp.sort_values(["__risk_rank__", "风险分值", "涉及金额", "证据笔数"], ascending=[False, False, False, False], kind="mergesort").head(5)
            desc = []
            for _, r in top_cs.iterrows():
                desc.append(
                    f"{safe_str(r.get('线索类型',''))}："
                    f"{safe_str(r.get('交易对方姓名','') or '对方缺失')}，"
                    f"{safe_str(r.get('规则编号',''))}，"
                    f"分值{_fmt_human_num(r.get('风险分值',0))}，"
                    f"证据{_fmt_human_num(r.get('证据笔数',0))}笔，"
                    f"涉及{_fmt_money_human(r.get('涉及金额',0))}，"
                    f"风险等级{safe_str(r.get('风险等级','')) or '未分级'}"
                )
            if desc:
                lines.append("重点项摘述：" + "；".join(desc) + "。")
    lines.append("")

    # 五、综合提示
    lines.append("五、综合提示")
    lines.append(
        "本报告为自动化归集与统计结果，主要用于线索梳理与辅助研判；涉及身份信息、权属状态、交易真实性、资金性质及通信语境等，应以原始材料及权威查询结果为准。"
        "对异常大额、频繁现金交易、对手净流入异常集中、非工作时间/深夜高频通信，以及关系风险线索命中项等情况，建议结合交易对手关系、资金去向、业务背景和通联背景开展进一步核查。"
    )
    lines.append("")

    outp = out_report_dir(person) / f"{person}-汇总报告.txt"
    outp.write_text("\n".join(lines), encoding="utf-8")
    return outp


# ========= GUI =========
def create_gui():
    root = tk.Tk()
    root.title("温岭纪委初核工具")
    root.minsize(820, 600)

    ttk.Label(root, text="温岭纪委初核工具", font=("仿宋", 18, "bold")).grid(row=0, column=0, columnspan=3, pady=(15, 0))
    ttk.Label(root, text="© 温岭纪委六室 单柳昊", font=("微软雅黑", 9)).grid(row=1, column=0, columnspan=3, pady=(0, 6))

    ttk.Label(root, text="工作目录:").grid(row=2, column=0, sticky="e", padx=8, pady=8)
    path_var = tk.StringVar(value=str(Path.home()))
    ttk.Entry(root, textvariable=path_var, width=60).grid(row=2, column=1, sticky="we", padx=5, pady=8)
    ttk.Button(root, text="浏览...", command=lambda: path_var.set(filedialog.askdirectory(title="选择工作目录") or path_var.get())).grid(row=2, column=2, padx=5, pady=8)

    log_box = tk.Text(root, width=96, height=18, state="disabled")
    log_box.grid(row=4, column=0, columnspan=3, padx=10, pady=(5, 10), sticky="nsew")
    root.columnconfigure(1, weight=1)
    root.rowconfigure(4, weight=1)

    tip = (
        f"tips1：通讯录文件需表头包含：姓名、职务、号码；\n"
        f"tips2：通信文件命名需以“-通信.xlsx（或 .xls）”结尾，将优先读取“本方姓名”列作为归属人员（无此列则回退使用文件名）；\n"
        f"tips3：通信标注/统计输出路径：分析结果/<人名>/通信/已标注、统计。\n"
    )
    log_box.config(state="normal"); log_box.insert("end", tip + "\n"); log_box.config(state="disabled")

    def log(msg: str):
        log_box.config(state="normal")
        log_box.insert("end", f"{_dt.datetime.now():%H:%M:%S}  {msg}\n")
        log_box.config(state="disabled")
        log_box.see("end")

    def clear_log() -> None:
        log_box.config(state="normal")
        log_box.delete("1.0", "end")
        log_box.config(state="disabled")

    def run(path: str):
        clear_log()
        STATE.OUT_DIR = _ensure_dir(Path(path).expanduser().resolve() / OUTPUT_ROOT_NAME)

        _orig_print = builtins.print
        builtins.print = lambda *a, **k: log(" ".join(map(str, a)))

        try:
            if LunarDate is None:
                print("⚠️ 未检测到 lunardate 库，农历节日判定将使用近似法（建议：pip install lunardate）")

            all_txn = merge_all_txn(path)
            if all_txn.empty:
                messagebox.showinfo("完成", "未找到可分析文件")
                return

            STATE.CLUE_SUMMARY_BY_PERSON = {}
            STATE.CLUE_DETAIL_BY_PERSON = {}

            for person, df_person in all_txn.groupby("查询对象", dropna=False):
                person_name = safe_str(person).strip() or "未知"
                print(f"--- 分析 {person_name} ---")
                analysis_txn(df_person)
                make_partner_summary(df_person)
                generate_suspicious_clues(df_person)

            export_all_clue_summary()

            root_path = Path(path).expanduser().resolve()
            for person, df_person in all_txn.groupby("查询对象", dropna=False):
                person_name = safe_str(person).strip() or "未知"
                rp = build_person_report_txt(root_path, person_name, df_person)
                print(f"✅ 已生成个人汇总报告：{rp}")

            messagebox.showinfo("完成", f"全部分析完成！结果在:\n{STATE.OUT_DIR}\n")
        except Exception as e:
            messagebox.showerror("错误", str(e))
        finally:
            builtins.print = _orig_print

    ttk.Button(
        root,
        text="开始分析",
        command=lambda: threading.Thread(target=run, args=(path_var.get().strip(),), daemon=True).start(),
        width=18,
    ).grid(row=3, column=1, pady=10)
    root.mainloop()


if __name__ == "__main__":
    create_gui()
