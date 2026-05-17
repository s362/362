import os
import glob
import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext
import threading
from datetime import datetime
import openpyxl

class ExcelSearchApp:
    def __init__(self, root):
        self.root = root
        self.root.title("温岭纪委excel查询小工具")
        self.root.geometry("650x550")
        
        # 开发者信息
        self.lbl_dev = tk.Label(root, text="© 温岭纪委六室 单柳昊", fg="gray")
        self.lbl_dev.pack(pady=5)

        # 文件夹选择
        self.frame_folder = tk.Frame(root)
        self.frame_folder.pack(pady=10, padx=20, fill="x")
        self.lbl_folder = tk.Label(self.frame_folder, text="查询文件夹:")
        self.lbl_folder.pack(side="left")
        self.entry_folder = tk.Entry(self.frame_folder)
        self.entry_folder.pack(side="left", fill="x", expand=True, padx=5)
        self.btn_browse = tk.Button(self.frame_folder, text="选择文件夹", command=self.browse_folder)
        self.btn_browse.pack(side="right")

        # 查询内容输入
        self.frame_query = tk.Frame(root)
        self.frame_query.pack(pady=10, padx=20, fill="x")
        self.lbl_query = tk.Label(self.frame_query, text="查询关键字:")
        self.lbl_query.pack(side="left")
        
        self.placeholder = "用空格分隔多个关键词"
        self.entry_query = tk.Entry(self.frame_query, fg="gray")
        self.entry_query.insert(0, self.placeholder)
        self.entry_query.pack(side="left", fill="x", expand=True, padx=5)
        
        self.entry_query.bind("<FocusIn>", self.on_focus_in)
        self.entry_query.bind("<FocusOut>", self.on_focus_out)

        # 运行按钮
        self.btn_run = tk.Button(root, text="开始查询并导出", bg="#4CAF50", fg="white", 
                                 font=("Arial", 12, "bold"), command=self.start_task)
        self.btn_run.pack(pady=15)

        # 状态日志显示
        self.log_area = scrolledtext.ScrolledText(root, height=15, state='disabled', bg="#f4f4f4")
        self.log_area.pack(pady=10, padx=20, fill="both", expand=True)

    def on_focus_in(self, event):
        if self.entry_query.get() == self.placeholder:
            self.entry_query.delete(0, tk.END)
            self.entry_query.config(fg="black")

    def on_focus_out(self, event):
        if not self.entry_query.get():
            self.entry_query.insert(0, self.placeholder)
            self.entry_query.config(fg="gray")

    def log(self, message):
        """记录日志并实时刷新UI"""
        self.log_area.config(state='normal')
        self.log_area.insert(tk.END, f"[{datetime.now().strftime('%H:%M:%S')}] {message}\n")
        self.log_area.see(tk.END)
        self.log_area.config(state='disabled')
        self.root.update_idletasks() # 强制刷新界面

    def browse_folder(self):
        folder_selected = filedialog.askdirectory()
        if folder_selected:
            self.entry_folder.delete(0, tk.END)
            self.entry_folder.insert(0, folder_selected)

    def start_task(self):
        input_dir = self.entry_folder.get().strip()
        query_str = self.entry_query.get().strip()

        if not input_dir or not os.path.exists(input_dir):
            messagebox.showerror("错误", "请选择有效的查询文件夹")
            return
        if not query_str or query_str == self.placeholder:
            messagebox.showerror("错误", "请输入查询关键字")
            return

        self.btn_run.config(state="disabled", text="正在处理中，请勿关闭...")
        thread = threading.Thread(target=self.run_logic, args=(input_dir, query_str))
        thread.daemon = True
        thread.start()

    def process_large_file(self, file_path, file_name, keywords, results):
        """流式引擎：专门处理 > 500MB 的超大文件，防内存崩溃并显示进度"""
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                headers = []
                matched_rows = {kw: [] for kw in keywords}
                row_count = 0

                # 逐行读取，不占用大量内存
                for row in ws.iter_rows(values_only=True):
                    row_count += 1
                    
                    # 提取表头
                    if row_count == 1:
                        headers = [str(c) if c is not None else f'未命名列_{i}' for i, c in enumerate(row)]
                        continue
                    
                    # 进度汇报：每 5000 行在日志中汇报一次
                    if row_count % 5000 == 0:
                        self.log(f"   ▶ {sheet_name} 已扫描到第 {row_count} 行...")

                    if not any(row):  # 过滤全空行
                        continue

                    # 转换为小写字符串用于搜索
                    row_str = [str(cell).lower() if cell is not None else "" for cell in row]
                    
                    for kw in keywords:
                        if any(kw.lower() in cell for cell in row_str):
                            matched_rows[kw].append([file_name, sheet_name] + list(row))
                
                # 将流式读取到的匹配数据，转换为 Pandas DataFrame 格式以便后续合并
                for kw in keywords:
                    if matched_rows[kw]:
                        col_names = ['来源文件', '来源Sheet'] + headers
                        
                        # 规避某些行长度超标导致报错的问题
                        max_len = max(len(r) for r in matched_rows[kw])
                        if max_len > len(col_names):
                            col_names += [f"附加列_{i}" for i in range(len(col_names), max_len)]
                            
                        # 截齐数据长度
                        cleaned_rows = [r + [None]*(len(col_names)-len(r)) for r in matched_rows[kw]]
                        df = pd.DataFrame(cleaned_rows, columns=col_names)
                        results[kw].append(df)
                        
            wb.close()
        except Exception as e:
            self.log(f"读取超大文件 {file_name} 时发生错误: {e}")

    def run_logic(self, input_dir, query_str):
        try:
            keywords = query_str.split()
            self.log(f"🚀 任务开始。关键词: {keywords}")
            
            output_file = os.path.join(input_dir, f"查询结果_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
            results = {kw: [] for kw in keywords}

            file_paths = glob.glob(os.path.join(input_dir, "**", "*.xls*"), recursive=True)
            total_files = len(file_paths)
            
            if total_files == 0:
                self.log("未发现 Excel 文件！")
                self.reset_btn()
                return

            for idx, file_path in enumerate(file_paths):
                file_name = os.path.basename(file_path)
                # 忽略临时文件和自身生成的结果文件
                if file_name.startswith('~$') or file_name == os.path.basename(output_file): 
                    continue
                
                # 获取文件大小 (MB)
                file_size_mb = os.path.getsize(file_path) / (1024 * 1024)
                
                self.log(f"({idx+1}/{total_files}) 正在读取: {file_name} [{file_size_mb:.1f} MB]")
                
                # ========== 引擎智能切换 ==========
                if file_size_mb > 500 and file_path.endswith(('.xlsx', '.xlsm')):
                    self.log(f"⚡ 触发大文件优化机制，切换为流式引擎...")
                    self.process_large_file(file_path, file_name, keywords, results)
                else:
                    # 常规文件 (<500MB) 使用 Pandas 极速引擎
                    try:
                        with pd.ExcelFile(file_path) as xls:
                            for sheet in xls.sheet_names:
                                df = pd.read_excel(xls, sheet_name=sheet)
                                if df.empty: continue
                                
                                df_str = df.astype(str).fillna('')
                                for kw in keywords:
                                    mask = df_str.apply(lambda x: x.str.contains(kw, case=False, na=False, regex=False)).any(axis=1)
                                    matched_df = df[mask].copy()
                                    if not matched_df.empty:
                                        matched_df.insert(0, '来源Sheet', sheet)
                                        matched_df.insert(0, '来源文件', file_name)
                                        results[kw].append(matched_df)
                    except Exception as e:
                        self.log(f"跳过文件 {file_name}: {e}")

            # ========== 写入最终结果 ==========
            self.log("正在生成结果报告并合并数据...")
            with pd.ExcelWriter(output_file, engine='xlsxwriter') as writer:
                for kw in keywords:
                    sheet_name = str(kw)[:31].replace('/', '_').replace('\\', '_')
                    if results[kw]:
                        pd.concat(results[kw], ignore_index=True).to_excel(writer, sheet_name=sheet_name, index=False)
                    else:
                        pd.DataFrame({'结果': ['未搜索到内容']}).to_excel(writer, sheet_name=sheet_name, index=False)

            self.log(f"✅ 完成！结果已存至该目录下。")
            messagebox.showinfo("查询完成", f"结果文件：\n{output_file}")

        except Exception as e:
            self.log(f"❌ 发生致命错误: {e}")
            messagebox.showerror("运行错误", f"发生错误：\n{e}")
        finally:
            self.reset_btn()

    def reset_btn(self):
        self.btn_run.config(state="normal", text="开始查询并导出")

if __name__ == "__main__":
    root = tk.Tk()
    app = ExcelSearchApp(root)
    root.mainloop()